"""Vẽ ``TaiLieu`` ra Word (.docx) bằng python-docx (REQ-EXP-03, BR-XF-03).

Bảng được dựng trực tiếp bằng phần tử XML vì API ô của python-docx rất chậm với
bảng hàng nghìn dòng. Mỗi hàng có ``cantSplit`` (không ngắt dòng qua trang), hàng
tiêu đề có ``tblHeader`` (lặp lại khi bảng sang trang).
"""
import io

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

from . import bo_cuc
from .bo_cuc import Bang, Doan, GiuCung
from .trinh_bay import KhoiKy

PHONG = "Times New Roman"
_KICH_THUOC_MM = {("A4", "doc"): (210, 297), ("A4", "ngang"): (297, 210), ("A3", "ngang"): (420, 297)}
_CAN = {"trai": "left", "giua": "center", "phai": "right"}
_KIEU_DOAN = {
    "ten_so": dict(co_chu=14, can="center", dam=True, truoc=6, sau=2),
    "giua": dict(co_chu=10.5, can="center", sau=4),
    "thuong": dict(co_chu=10),
    "nghieng": dict(co_chu=10, nghieng=True),
    "muc": dict(co_chu=11, dam=True, truoc=8, sau=3),
    "nho": dict(co_chu=8.5),
    "dam": dict(co_chu=10, dam=True, truoc=4),
    "phai": dict(co_chu=10, can="right"),
}


def _mm_sang_twip(gia_tri):
    return int(round(gia_tri * 1440 / 25.4))


def _phan_tu(ten, **thuoc_tinh):
    pt = OxmlElement(ten)
    for k, v in thuoc_tinh.items():
        pt.set(qn(f"w:{k}"), str(v))
    return pt


def _run(chu, co_chu, dam=False, nghieng=False, gach_chan=False):
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    if dam:
        rpr.append(OxmlElement("w:b"))
    if nghieng:
        rpr.append(OxmlElement("w:i"))
    if gach_chan:
        rpr.append(_phan_tu("w:u", val="single"))
    rpr.append(_phan_tu("w:sz", val=int(co_chu * 2)))
    rpr.append(_phan_tu("w:szCs", val=int(co_chu * 2)))
    r.append(rpr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = chu
    r.append(t)
    return r


def _p(dong, co_chu, can="left", dam=False, nghieng=False, nhan_manh="", giu_voi_sau=False, co_chu_dong_sau=None):
    """Đoạn gồm nhiều dòng (ngắt bằng w:br); ``nhan_manh`` in đậm, gạch chân."""
    p = OxmlElement("w:p")
    ppr = OxmlElement("w:pPr")
    if giu_voi_sau:
        ppr.append(OxmlElement("w:keepNext"))
    ppr.append(_phan_tu("w:spacing", before=0, after=0))
    ppr.append(_phan_tu("w:jc", val=can))
    p.append(ppr)
    for i, chu in enumerate(dong):
        if i:
            r = OxmlElement("w:r")
            r.append(OxmlElement("w:br"))
            p.append(r)
            if co_chu_dong_sau:
                co_chu = co_chu_dong_sau
        if nhan_manh and nhan_manh in chu:
            truoc, sau = chu.split(nhan_manh, 1)
            if truoc:
                p.append(_run(truoc, co_chu, dam, nghieng))
            p.append(_run(nhan_manh, co_chu, True, nghieng, gach_chan=True))
            if sau:
                p.append(_run(sau, co_chu, dam, nghieng))
        else:
            p.append(_run(chu, co_chu, dam, nghieng))
    return p


def _tc(rong_twip, phan_tu_p, gop=1, nen=None):
    tc = OxmlElement("w:tc")
    tcpr = OxmlElement("w:tcPr")
    tcpr.append(_phan_tu("w:tcW", w=rong_twip, type="dxa"))
    if gop > 1:
        tcpr.append(_phan_tu("w:gridSpan", val=gop))
    if nen:
        tcpr.append(_phan_tu("w:shd", val="clear", color="auto", fill=nen))
    tc.append(tcpr)
    tc.append(phan_tu_p)
    return tc


def _tr(cac_tc, tieu_de=False):
    tr = OxmlElement("w:tr")
    trpr = OxmlElement("w:trPr")
    trpr.append(OxmlElement("w:cantSplit"))  # Không ngắt một dòng qua hai trang (BR-IN-09).
    if tieu_de:
        trpr.append(OxmlElement("w:tblHeader"))  # Lặp hàng tiêu đề khi bảng sang trang (BR-IN-04).
    tr.append(trpr)
    for tc in cac_tc:
        tr.append(tc)
    return tr


def _vien(khung):
    vien = OxmlElement("w:tblBorders")
    for canh in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if khung:
            vien.append(_phan_tu(f"w:{canh}", val="single", sz=4, space=0, color="000000"))
        else:
            vien.append(_phan_tu(f"w:{canh}", val="nil"))
    return vien


def _tbl(rong_cot_twip, khung):
    tbl = OxmlElement("w:tbl")
    tblpr = OxmlElement("w:tblPr")
    tblpr.append(_phan_tu("w:tblW", w=sum(rong_cot_twip), type="dxa"))
    tblpr.append(_vien(khung))
    tblpr.append(_phan_tu("w:tblLayout", type="fixed"))
    lam = OxmlElement("w:tblCellMar")
    for canh, gia_tri in (("left", 50), ("right", 50)):
        lam.append(_phan_tu(f"w:{canh}", w=gia_tri, type="dxa"))
    tblpr.append(lam)
    tbl.append(tblpr)
    luoi = OxmlElement("w:tblGrid")
    for w in rong_cot_twip:
        luoi.append(_phan_tu("w:gridCol", w=w))
    tbl.append(luoi)
    return tbl


def _bang(bang, rong_mm, giu_cung=False):
    tong = sum(c.ty_le for c in bang.cot)
    rong = [_mm_sang_twip(rong_mm * c.ty_le / tong) for c in bang.cot]
    tbl = _tbl(rong, bang.khung)
    if bang.nhom_tieu_de:
        cac_tc, cot = [], 0
        for nhan, so_cot in bang.nhom_tieu_de:
            cac_tc.append(_tc(sum(rong[cot:cot + so_cot]), _p([nhan], bang.co_chu, "center", dam=True),
                              so_cot, "D9D9D9"))
            cot += so_cot
        tbl.append(_tr(cac_tc, tieu_de=True))
    if bang.co_tieu_de:
        tbl.append(_tr([_tc(w, _p(c.tieu_de.split("\n"), bang.co_chu - 0.5, "center", dam=True), nen="D9D9D9")
                        for w, c in zip(rong, bang.cot)], tieu_de=True))
    for hang in bang.dong:
        tbl.append(_tr([
            _tc(w, _p(o.dong, bang.co_chu, _CAN[o.can], o.dam, o.nghieng, o.nhan_manh, giu_voi_sau=giu_cung,
                      co_chu_dong_sau=o.co_chu_dong_sau))
            for w, o in zip(rong, hang)
        ]))
    return tbl


def _doan(doan, giu_voi_sau=False):
    k = dict(_KIEU_DOAN[doan.kieu])
    p = _p(doan.chu.split("\n"), k["co_chu"], k.get("can", "left"), k.get("dam", False), k.get("nghieng", False),
           giu_voi_sau=giu_voi_sau)
    ppr = p.find(qn("w:pPr"))
    ppr.find(qn("w:spacing")).set(qn("w:before"), str(int(k.get("truoc", 0) * 20)))
    ppr.find(qn("w:spacing")).set(qn("w:after"), str(int(k.get("sau", 0) * 20)))
    return p


def _khoi_ky(khoi, rong_mm):
    n = len(khoi.o)
    rong = [_mm_sang_twip(rong_mm / n)] * n
    tbl = _tbl(rong, khung=False)
    for dong, co_chu, dam, nghieng in (
        ([o.chuc_danh for o in khoi.o], 10, True, False),
        (["(Ký, ghi rõ họ tên)"] * n, 9, False, True),
        ([""] * n, 10, False, False), ([""] * n, 10, False, False), ([""] * n, 10, False, False),
        ([o.ten or "..............................." for o in khoi.o], 10, False, False),
    ):
        tbl.append(_tr([_tc(w, _p([chu], co_chu, "center", dam, nghieng, giu_voi_sau=True))
                        for w, chu in zip(rong, dong)]))
    return [_p([khoi.dong_ngay], 10, "right", nghieng=True, giu_voi_sau=True), tbl]


def _them(than, khoi, rong_mm, giu_cung=False):
    if isinstance(khoi, Doan):
        than.append(_doan(khoi, giu_voi_sau=giu_cung))
    elif isinstance(khoi, Bang):
        than.append(_bang(khoi, rong_mm, giu_cung))
        than.append(_p([""], 4))  # Word cần một đoạn ngăn giữa hai bảng liền nhau.
    elif isinstance(khoi, KhoiKy):
        for pt in _khoi_ky(khoi, rong_mm):
            than.append(pt)
    elif isinstance(khoi, GiuCung):
        for k in khoi.khoi:
            _them(than, k, rong_mm, giu_cung=True)
    else:
        raise TypeError(khoi)


def _truong(p, lenh):
    """Trường PAGE/NUMPAGES của Word."""
    fld = _phan_tu("w:fldSimple", instr=lenh)
    fld.append(_run("1", 8))
    p._p.append(fld)


def _thiet_lap_section(section, phan, tai_lieu, dau_tien):
    rong, cao = _KICH_THUOC_MM[(phan.kho, phan.huong)]
    section.orientation = WD_ORIENT.LANDSCAPE if phan.huong == "ngang" else WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = Mm(rong), Mm(cao)
    section.left_margin, section.right_margin = Mm(bo_cuc.LE_TRAI_MM), Mm(bo_cuc.LE_PHAI_MM)
    section.top_margin, section.bottom_margin = Mm(bo_cuc.LE_TREN_MM), Mm(bo_cuc.LE_DUOI_MM)
    section.header_distance = section.footer_distance = Mm(6)
    if not dau_tien:
        # Đầu/chân trang kế thừa section trước; trang đầu của section sau vẫn có đầu trang rút gọn.
        section.different_first_page_header_footer = False
        return
    section.different_first_page_header_footer = True  # Trang 1 có khối đầu sổ, không cần đầu trang rút gọn.
    p = section.header.paragraphs[0]
    p.text = tai_lieu.dau_trang_rut_gon
    p.runs[0].font.size = Pt(8)
    for chan in (section.footer, section.first_page_footer):
        p = chan.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(f"{tai_lieu.chan_trang} · Trang ").font.size = Pt(8)
        _truong(p, "PAGE")
        p.add_run("/").font.size = Pt(8)
        _truong(p, "NUMPAGES")


def tao_docx(tai_lieu):
    doc = Document()
    kieu = doc.styles["Normal"]
    kieu.font.name = PHONG
    kieu.font.size = Pt(10)
    rfonts = kieu.element.get_or_add_rPr().get_or_add_rFonts()
    for thuoc_tinh in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(thuoc_tinh), PHONG)
    kieu.paragraph_format.space_after = Pt(0)
    kieu.paragraph_format.space_before = Pt(0)

    than = doc.element.body
    for p in than.findall(qn("w:p")):  # Bỏ đoạn trống mặc định của tài liệu mới.
        than.remove(p)
    for i, phan in enumerate(tai_lieu.phan):
        # add_section kết thúc section trước bằng một đoạn chứa bản sao sectPr của nó; sectPr
        # ở cuối thân trở thành section mới. Nội dung luôn chèn ngay trước sectPr cuối thân.
        section = doc.sections[0] if i == 0 else doc.add_section(WD_SECTION.NEW_PAGE)
        _thiet_lap_section(section, phan, tai_lieu, i == 0)
        sect_pr = than.find(qn("w:sectPr"))
        tam = []
        for khoi in phan.khoi:
            _them(tam, khoi, bo_cuc.rong_vung_in(phan.kho, phan.huong))
        for pt in tam:
            sect_pr.addprevious(pt)
    cp = doc.core_properties
    cp.title = tai_lieu.tieu_de
    cp.author = tai_lieu.tac_gia
    cp.subject = tai_lieu.chu_de
    cp.comments = "Mẫu nghiệp vụ nội bộ — bản xuất để biên tập. Mã trích xuất: " + tai_lieu.chu_de
    cp.keywords = tai_lieu.chu_de
    bo_dem = io.BytesIO()
    doc.save(bo_dem)
    return bo_dem.getvalue()

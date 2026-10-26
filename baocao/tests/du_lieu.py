"""Dựng bộ dữ liệu DS-01 (test_plan.md §5.1) và các tiện ích kiểm thử."""
import datetime
import io
import zipfile

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.utils import timezone
from openpyxl import load_workbook

from ravao.models import CanBo, Cong, DonVi, Khach, Luot, PhanCongCong, PhuongTien

QUYEN_DAY_DU = [
    "xem_so_cbcs", "xuat_so_cbcs", "xem_so_khach", "xuat_so_khach", "xem_so_phuong_tien",
    "xuat_so_phuong_tien", "xem_so_tong_hop", "xuat_so_tong_hop", "xem_moi_cong",
]


def t(chuoi):
    """'HH:MM dd/mm' (năm 2026) hoặc 'HH:MM dd/mm/yyyy' theo giờ địa phương."""
    if chuoi.count("/") == 1:
        chuoi += "/2026"
    return timezone.make_aware(datetime.datetime.strptime(chuoi, "%H:%M %d/%m/%Y"))


def tao_nguoi_dung(ten, ho="", ten_rieng="", quyen=()):
    nguoi = get_user_model().objects.create_user(ten, password="mat-khau-thu", last_name=ho, first_name=ten_rieng)
    if quyen:
        nguoi.user_permissions.set(Permission.objects.filter(content_type__app_label="ravao", codename__in=quyen))
    return get_user_model().objects.get(pk=nguoi.pk)  # Xoá bộ nhớ đệm quyền.


class DS01:
    """Các đối tượng của DS-01, truy cập bằng tên lượt: ``ds.l["C3"]``."""

    def __init__(self):
        self.tb01 = tao_nguoi_dung("trucban01", "Nguyễn Văn", "X", QUYEN_DAY_DU)
        self.tb02 = tao_nguoi_dung("trucban02", "Lê Thị", "Y", QUYEN_DAY_DU)
        self.tb03 = tao_nguoi_dung("trucban03", "Trần Văn", "Z", QUYEN_DAY_DU)
        self.chihuy = tao_nguoi_dung(
            "chihuy01", "Phan Văn", "H", QUYEN_DAY_DU + ["xem_ban_ghi_huy", "xem_du_lieu_ca_nhan_day_du"]
        )
        self.c1 = Cong.objects.create(ma="C1", ten="Cổng 1", thu_tu=1)
        self.c2 = Cong.objects.create(ma="C2", ten="Cổng 2", thu_tu=2)
        tham_muu = DonVi.objects.create(ten="Phòng Tham mưu")
        hau_can = DonVi.objects.create(ten="Phòng Hậu cần")
        ky_thuat = DonVi.objects.create(ten="Phòng Kỹ thuật")
        cb = {}
        for ma, ten, dv in [
            ("CB001", "Nguyễn Văn An", tham_muu), ("CB002", "Trần Thị Bình", hau_can),
            ("CB003", "Lê Văn Cường", ky_thuat), ("CB004", "Phạm Thị Dung", hau_can),
            ("CB005", "Hoàng Văn Đức", ky_thuat), ("CB006", "Vũ Thị Giang", tham_muu),
            ("CB007", "Đặng Quốc Hưng", tham_muu),
        ]:
            cb[ma] = CanBo.objects.create(ma=ma, ho_ten=ten, don_vi=dv)
        self.cb = cb

        lam = Khach.objects.create(ho_ten="Bùi Văn Lâm", loai_giay_to="CCCD", so_giay_to="001090000111",
                                   co_quan="Công ty TNHH An Phát", so_dien_thoai="0912345678")
        lam2 = Khach.objects.create(ho_ten="Bùi Văn Lâm", loai_giay_to="CCCD", so_giay_to="001 090 000 111",
                                    co_quan="Công ty TNHH An Phát")
        ngoc = Khach.objects.create(ho_ten="Mai Thị Ngọc", loai_giay_to="CCCD", so_giay_to="001190000222")
        phuc = Khach.objects.create(ho_ten="Tạ Văn Phúc", loai_giay_to="CCCD", so_giay_to="001090000333")
        quang = Khach.objects.create(ho_ten="Ông Quang")
        sen = Khach.objects.create(ho_ten="Lý Thị Sen", loai_giay_to="CCCD", so_giay_to="001190000444")

        p_co_quan = PhuongTien.objects.create(bien_so="30A-123.45", loai="Ô tô", nguon="CO_QUAN")
        p_co_quan_2 = PhuongTien.objects.create(bien_so="30A12345", loai="Ô tô", nguon="CO_QUAN")
        p_khach = PhuongTien.objects.create(bien_so="29B-678.90", loai="Ô tô", nguon="KHACH", nguon_khach=lam)
        p_xe_may = PhuongTien.objects.create(bien_so="29X1-111.11", loai="Xe máy", nguon="CBCS",
                                             nguon_can_bo=cb["CB004"])
        p_khong_bien = PhuongTien.objects.create(ma_tam="KB-001", loai="Xe chuyên dùng", nguon="KHAC",
                                                 nguon_ghi_chu="Đơn vị thi công")
        self.pt = {"P1": p_co_quan, "P2": p_co_quan_2, "P3": p_khach, "P4": p_xe_may, "P5": p_khong_bien}

        self.l = {}
        L = self._luot
        L("C0", "CBCS", "RV", "09:00 24/10", "10:00 24/10", can_bo=cb["CB001"], ly_do="Công tác")
        L("C3", "CBCS", "RV", "23:30 24/10", "01:15 25/10", can_bo=cb["CB002"], ly_do="Việc riêng có phép",
          mo_nguoi=self.tb03)
        L("C8", "CBCS", "VR", "07:30 25/10", "17:15 25/10", can_bo=cb["CB007"], dong_nguoi=self.tb02)
        L("C1", "CBCS", "RV", "08:00 25/10", "10:30 25/10", can_bo=cb["CB001"], ly_do="Công tác",
          phuong_tien_di_kem=p_co_quan)
        L("C6", "CBCS", "RV", "09:00 25/10", "09:14 25/10", can_bo=cb["CB005"], ly_do="Công tác",
          dong_ghi="09:41 25/10")
        L("C7", "CBCS", "RV", "11:00 25/10", None, can_bo=cb["CB006"])
        L("C2", "CBCS", "RV", "14:00 25/10", "15:00 25/10", can_bo=cb["CB001"], ly_do="Công tác",
          mo_cong=self.c2, dong_cong=self.c2, phuong_tien_di_kem=p_co_quan_2)
        L("C5", "CBCS", "RV", "16:00 25/10", None, can_bo=cb["CB004"], ly_do="Việc riêng có phép",
          han="18:00 25/10", phuong_tien_di_kem=p_xe_may)
        L("C4", "CBCS", "RV", "22:45 25/10", "00:20 26/10", can_bo=cb["CB003"], ly_do="Công tác", mo_cong=self.c2)

        L("K1", "KHACH", "VR", "08:30 25/10", "09:45 25/10", khach=lam, so_the="K01", phuong_tien_di_kem=p_khach,
          muc_dich="Liên hệ công tác", nguoi_tiep="Nguyễn Văn An (CB001)")
        L("K3", "KHACH", "VR", "10:00 25/10", "11:00 25/10", khach=ngoc, so_the="K03", ma_doan="ĐK-01",
          muc_dich="Làm việc theo kế hoạch", nguoi_tiep="Phòng Tham mưu")
        L("K4", "KHACH", "VR", "10:00 25/10", "11:05 25/10", khach=phuc, so_the="K04", ma_doan="ĐK-01",
          muc_dich="Làm việc theo kế hoạch", nguoi_tiep="Phòng Tham mưu", dong_ghi="12:10 25/10")
        L("K2", "KHACH", "VR", "14:00 25/10", "14:30 25/10", khach=lam2, so_the="K02",
          muc_dich="Nộp hồ sơ", nguoi_tiep="Phòng Hậu cần")
        L("K5", "KHACH", "VR", "15:00 25/10", None, khach=quang, so_the="K05", han="17:30 25/10",
          muc_dich="Giao hàng", nguoi_tiep="Phòng Hậu cần")
        L("K6", "KHACH", "VR", "17:30 25/10", "00:10 26/10", khach=sen, so_the="K06",
          muc_dich="Thăm thân", nguoi_tiep="Phòng Kỹ thuật")

        L("P1", "PT", "RV", "08:00 25/10", "10:30 25/10", phuong_tien=p_co_quan, lenh_dieu_xe="LĐX-101",
          nguoi_dieu_khien_can_bo=cb["CB001"], ly_do="Công tác")
        L("P3", "PT", "VR", "08:28 25/10", "09:50 25/10", phuong_tien=p_khach, nguoi_dieu_khien="Bùi Văn Lâm")
        L("P5", "PT", "VR", "10:00 25/10", "10:30 25/10", phuong_tien=p_khong_bien, nguoi_dieu_khien="Nguyễn Văn Tâm")
        L("P2", "PT", "RV", "14:00 25/10", "15:00 25/10", phuong_tien=p_co_quan_2, lenh_dieu_xe="LĐX-102",
          nguoi_dieu_khien_can_bo=cb["CB001"], mo_cong=self.c2, dong_cong=self.c2)
        L("P4", "PT", "RV", "16:00 25/10", None, phuong_tien=p_xe_may, nguoi_dieu_khien_can_bo=cb["CB004"])

        self.l["C7"].huy("Nhập nhầm người", self.chihuy)
        Luot.objects.filter(pk=self.l["C7"].pk).update(huy_luc=t("11:10 25/10"))
        dc = self.l["C6"].dinh_chinh("dong_thoi_diem", t("09:40 25/10"), "Nhập nhầm giờ", self.tb02)
        dc.thoi_diem = t("20:00 25/10")
        dc.save()

    def _luot(self, ten, nhom, chieu, mo, dong, mo_cong=None, dong_cong=None, mo_nguoi=None, dong_nguoi=None,
              dong_ghi=None, han=None, **kw):
        mo_tg = t(mo)
        dong_tg = t(dong) if dong else None
        luot = Luot.objects.create(
            nhom=nhom, chieu=chieu,
            mo_thoi_diem=mo_tg, mo_cong=mo_cong or self.c1, mo_ghi_nhan_luc=mo_tg + datetime.timedelta(minutes=1),
            mo_nguoi_ghi=mo_nguoi or self.tb01,
            dong_thoi_diem=dong_tg,
            dong_cong=(dong_cong or self.c1) if dong else None,
            dong_ghi_nhan_luc=(t(dong_ghi) if dong_ghi else dong_tg + datetime.timedelta(minutes=1)) if dong else None,
            dong_nguoi_ghi=(dong_nguoi or self.tb01) if dong else None,
            han_du_kien=t(han) if han else None,
            **kw,
        )
        self.l[ten] = luot
        return luot

    def ten(self, ma_luot_id):
        return {v.pk: k for k, v in self.l.items()}[ma_luot_id]


def phan_cong(nguoi, *cong):
    for c in cong:
        PhanCongCong.objects.create(nguoi_dung=nguoi, cong=c)


# ---------------------------------------------------------------------------
# Đọc tệp Excel
# ---------------------------------------------------------------------------


def mo_workbook(noi_dung):
    return load_workbook(io.BytesIO(noi_dung))


def hang_tieu_de(ws):
    for hang in range(1, 40):
        if ws.cell(hang, 1).value == "STT":
            return hang
    raise AssertionError("Không thấy hàng tiêu đề bảng")


def doc_bang(ws, hang_dau=None):
    """Trả về (tiêu đề, [dict theo tiêu đề]) của bảng bắt đầu ở hàng tiêu đề."""
    hang_dau = hang_dau or hang_tieu_de(ws)
    tieu_de = [c.value for c in ws[hang_dau] if c.value is not None]
    dong = []
    for hang in ws.iter_rows(min_row=hang_dau + 1, max_col=len(tieu_de)):
        if hang[0].value is None or (isinstance(hang[0].value, str) and hang[0].value.startswith("Không có")):
            break
        dong.append({tieu_de[i]: o.value for i, o in enumerate(hang)})
    return tieu_de, dong


def xml_trong_tep(noi_dung):
    with zipfile.ZipFile(io.BytesIO(noi_dung)) as z:
        return {ten: z.read(ten).decode("utf-8") for ten in z.namelist() if ten.endswith(".xml")}

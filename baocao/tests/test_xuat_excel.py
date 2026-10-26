"""Xuất Excel qua giao diện HTTP: nội dung tệp, bố cục, an toàn, quyền, lịch sử xuất."""
import datetime
import hashlib
import time
import unicodedata
from unittest import mock

from django.test import TestCase, TransactionTestCase
from django.urls import reverse
from openpyxl.utils import get_column_letter

from ravao.models import CanBo, DinhChinh, DonVi, Khach, Luot

from .. import excel
from ..models import CauHinhSo, LichSuXuat
from .du_lieu import DS01, doc_bang, hang_tieu_de, mo_workbook, phan_cong, t, tao_nguoi_dung, xml_trong_tep

S1 = {"tu": "2026-10-25T00:00", "den": "2026-10-26T00:00"}
GIO_S1 = t("09:00 27/10")


def url(so):
    return reverse("baocao:xuat_excel", args=[so])


class CoSo(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.ds = DS01()
        CauHinhSo.objects.create(co_quan_chu_quan="Cơ quan thử nghiệm", ten_don_vi="Đơn vị thử nghiệm",
                                 dia_danh="Hà Nội")

    def setUp(self):
        dong_ho = mock.patch("baocao.dong_ho.hien_tai", return_value=GIO_S1)
        dong_ho.start()
        self.addCleanup(dong_ho.stop)

    def xuat(self, so="cbcs", nguoi=None, **tham_so):
        self.client.force_login(nguoi or self.ds.tb01)
        return self.client.get(url(so), {**S1, **tham_so})

    def xuat_ok(self, so="cbcs", nguoi=None, **tham_so):
        phan_hoi = self.xuat(so, nguoi, **tham_so)
        self.assertEqual(phan_hoi.status_code, 200, phan_hoi.content[:500])
        return phan_hoi, mo_workbook(phan_hoi.content)


class NoiDungSoChiTietTest(CoSo):
    def test_so01_tieu_de_thong_tin_va_dong(self):
        phan_hoi, wb = self.xuat_ok()
        self.assertEqual(wb.sheetnames, ["SO01", "Dinh_chinh", "Thong_tin_trich_xuat"])
        ws = wb["SO01"]
        dau = [ws.cell(h, 1).value for h in range(1, 9)]
        self.assertEqual(dau[0], "CƠ QUAN THỬ NGHIỆM")
        self.assertEqual(dau[2], "SỔ THEO DÕI CBCS RA/VÀO")
        self.assertEqual(dau[3], "Từ 00:00 ngày 25/10/2026 đến hết ngày 25/10/2026")
        self.assertEqual(dau[5], "Số liệu tính đến: 00:00:00 ngày 26/10/2026")
        ma = phan_hoi["X-Ma-Trich-Xuat"]
        self.assertRegex(ma, r"^SO01-20261027-\d{4}$")
        self.assertIn(f"Thời điểm trích xuất: 09:00:00 ngày 27/10/2026 (giờ máy chủ)    Mã trích xuất: {ma}", dau[6])
        self.assertEqual(dau[7], "Người xuất: Nguyễn Văn X (tài khoản: trucban01)    Người lập sổ: Nguyễn Văn X")

        _, dong = doc_bang(ws)
        self.assertEqual([d["Mã lượt"] for d in dong], [self.ds.l[k].pk for k in ["C3", "C8", "C1", "C6", "C2", "C5", "C4"]])
        self.assertEqual([d["STT"] for d in dong], [1, 2, 3, 4, 5, 6, 7])
        c6 = dong[3]
        self.assertEqual((c6["Vào – thời điểm"], c6["Ghi chú"]), (datetime.datetime(2026, 10, 25, 9, 40), "ĐC(1)"))
        tong_hop = {ws.cell(h, 2).value: ws.cell(h, 6).value for h in range(1, ws.max_row + 1) if ws.cell(h, 2).value}
        self.assertEqual(tong_hop["Số lượt trong sổ"], 7)
        self.assertEqual(tong_hop["— Từ kỳ trước"], 1)
        self.assertEqual(tong_hop["Chưa hoàn tất tại thời điểm chốt"], 2)
        self.assertEqual(tong_hop["Số CBCS duy nhất"], 6)
        self.assertIn("Cân đối: Từ kỳ trước + Phát sinh = Hoàn tất + Chưa hoàn tất (1 + 6 = 5 + 2)", tong_hop)
        self.assertEqual(phan_hoi["Content-Disposition"],
                         f'attachment; filename="SO01_CBCS_20261025-20261025_{ma}.xlsx"')

    def test_qua_nua_dem_trong_tep(self):
        _, wb = self.xuat_ok()
        dong = {d["Mã lượt"]: d for d in doc_bang(wb["SO01"])[1]}
        c3, c4, c5 = (dong[self.ds.l[k].pk] for k in ("C3", "C4", "C5"))
        self.assertEqual((c3["Ra – thời điểm"], c3["Ra – chú thích kỳ"]),
                         (datetime.datetime(2026, 10, 24, 23, 30), "(từ kỳ trước)"))
        self.assertEqual(c3["Vào – thời điểm"], datetime.datetime(2026, 10, 25, 1, 15))
        self.assertEqual((c4["Vào – thời điểm"], c4["Vào – chú thích kỳ"], c4["Trạng thái"]),
                         (datetime.datetime(2026, 10, 26, 0, 20), "(sau kỳ)", "Hoàn tất sau kỳ"))
        self.assertEqual(c4["Thời lượng (giờ:phút)"], datetime.timedelta(hours=1, minutes=35))
        self.assertEqual((c5["Vào – thời điểm"], c5["Thời lượng (giờ:phút)"], c5["Là thời lượng tối thiểu"],
                          c5["Trạng thái"]), (None, datetime.timedelta(hours=8), "Có", "Chưa về – Quá hạn"))

    def test_dinh_dang_ngay_gio_nhat_quan_va_kieu_so(self):
        _, wb = self.xuat_ok()
        ws = wb["SO01"]
        h = hang_tieu_de(ws)
        tieu_de = [c.value for c in ws[h]]
        for hang in ws.iter_rows(min_row=h + 1, max_row=h + 7):
            for ten, o in zip(tieu_de, hang):
                if ten and (ten.endswith("thời điểm") or ten.endswith("lúc") or ten == "Hạn về") and o.value:
                    self.assertIsInstance(o.value, datetime.datetime, ten)
                    self.assertEqual(o.number_format, "dd/mm/yyyy hh:mm:ss", ten)
                if ten == "Thời lượng (giờ:phút)":
                    self.assertIn(o.number_format, ("[h]:mm", '"≥ "[h]:mm'))
        # Lượt chưa hoàn tất: vẫn là giá trị số nhưng hiển thị "≥ 8:00" (BR-ND-05).
        c5 = next(hang for hang in ws.iter_rows(min_row=h + 1) if hang[0].value == 6)
        o = c5[tieu_de.index("Thời lượng (giờ:phút)")]
        self.assertEqual((o.value, o.number_format), (datetime.timedelta(hours=8), '"≥ "[h]:mm'))

    def test_phu_luc_va_thong_tin_trich_xuat(self):
        phan_hoi, wb = self.xuat_ok()
        _, dc = doc_bang(wb["Dinh_chinh"], 4)
        self.assertEqual(dc, [{
            "(n)": 1, "STT dòng": 4, "Đối tượng": "CB005 Hoàng Văn Đức", "Trường": "Thời điểm sự kiện đóng",
            "Giá trị cũ": "09:14 25/10/2026", "Giá trị mới": "09:40 25/10/2026", "Lý do": "Nhập nhầm giờ",
            "Người đính chính": "Lê Thị Y (trucban02)", "Thời điểm đính chính": datetime.datetime(2026, 10, 25, 20, 0),
        }])
        tt = {r[0].value: r[1].value for r in wb["Thong_tin_trich_xuat"].iter_rows(max_col=2) if r[0].value}
        self.assertEqual(tt["Mã trích xuất"], phan_hoi["X-Ma-Trich-Xuat"])
        self.assertEqual(tt["Kỳ — đến (không gồm)"], datetime.datetime(2026, 10, 26))
        self.assertEqual(tt["Người lập sổ"], "Nguyễn Văn X")
        self.assertEqual(tt["Số dòng dữ liệu"], 7)

    def test_nguoi_lap_so(self):
        _, wb = self.xuat_ok(nguoi_lap_so="trucban03")
        self.assertTrue(wb["SO01"].cell(8, 1).value.endswith("Người lập sổ: Trần Văn Z"))
        _, wb = self.xuat_ok(nguoi_lap_so="__trong__")
        self.assertTrue(wb["SO01"].cell(8, 1).value.endswith("(để trống — ký tay)"))
        self.assertEqual(self.xuat(nguoi_lap_so="khong-ton-tai").status_code, 400)

    def test_bo_loc_ghi_len_so(self):
        _, wb = self.xuat_ok(cong=self.ds.c2.pk, trang_thai="CHUA_HOAN_TAT")
        ws = wb["SO01"]
        self.assertEqual(ws.cell(5, 1).value,
                         "Điểm kiểm soát: Cổng 2    Điều kiện lọc: Trạng thái: Chưa hoàn tất")
        self.assertEqual(ws.cell(9, 1).value, "Sổ đã lọc — không phải toàn bộ dữ liệu của kỳ.")
        self.assertEqual([d["Mã lượt"] for d in doc_bang(ws)[1]], [self.ds.l["C4"].pk])

    def test_so02_va_so03(self):
        _, wb = self.xuat_ok("khach")
        dong = doc_bang(wb["SO02"])[1]
        self.assertEqual([d["Số thẻ"] for d in dong], ["K01", "K03", "K04", "K02", "K05", "K06"])
        self.assertEqual(dong[2]["Ghi chú"], "GBS")
        self.assertEqual(dong[4]["Số giấy tờ"], "Không xuất trình")
        _, wb = self.xuat_ok("phuong-tien")
        dong = doc_bang(wb["SO03"])[1]
        self.assertEqual([d["Biển số"] for d in dong],
                         ["30A-123.45", "29B-678.90", "Không biển — KB-001", "30A12345", "29X1-111.11"])
        self.assertEqual(dong[0]["Người điều khiển"], "Nguyễn Văn An (CB001)")
        self.assertEqual(dong[4]["Nguồn"], "CBCS — Phạm Thị Dung")


class ThietLapInTest(CoSo):
    def test_kho_giay_huong_vung_in_tieu_de_lap(self):
        _, wb = self.xuat_ok()
        ws = wb["SO01"]
        h = hang_tieu_de(ws)
        tieu_de = [c.value for c in ws[h] if c.value]
        so_cot = excel.so_cot_in(excel.cac_cot("CBCS", False))
        # Vùng in chỉ gồm cột in; cột chỉ dành cho Excel nằm bên phải, ngoài vùng in.
        self.assertEqual(tieu_de[so_cot - 1], "Ghi chú")
        self.assertEqual(set(tieu_de[so_cot:]), {
            "Cấp bậc", "Chức vụ", "Là thời lượng tối thiểu", "Nội dung", "Hạn về", "Ghi nhận Ra lúc",
            "Ghi nhận Vào lúc", "Mã trạng thái", "Mã lượt"})
        self.assertEqual(int(ws.page_setup.paperSize), int(ws.PAPERSIZE_A4))
        self.assertEqual(ws.page_setup.orientation, "landscape")
        self.assertTrue(ws.sheet_properties.pageSetUpPr.fitToPage)
        self.assertEqual((ws.page_setup.fitToWidth, ws.page_setup.fitToHeight), (1, 0))
        self.assertEqual(ws.print_title_rows, f"${h}:${h}")
        self.assertRegex(ws.print_area, rf"^'SO01'!\$A\$1:\${get_column_letter(so_cot)}\$\d+$")
        self.assertGreater(int(ws.print_area.rsplit("$", 1)[1]), h + 7)  # gồm cả khối tổng hợp và khối ký
        self.assertEqual(ws.auto_filter.ref, f"A{h}:{get_column_letter(len(tieu_de))}{h + 7}")
        self.assertEqual(ws.freeze_panes, f"A{h + 1}")
        self.assertIn("Trang &P/&N", ws.oddFooter.center.text)
        self.assertAlmostEqual(ws.page_margins.left, 20 / 25.4)
        for chu, cot in ws.column_dimensions.items():
            self.assertGreaterEqual(cot.width, 5, chu)
            self.assertLessEqual(cot.width, 62, chu)

    def test_so_tong_hop_kho_doc(self):
        _, wb = self.xuat_ok("tong-hop")
        self.assertEqual(wb.sheetnames, ["A_Chi_tieu", "B_Theo_ngay", "C_Theo_cong", "D_Chua_hoan_tat",
                                         "E_Dinh_chinh_huy", "Thong_tin_trich_xuat"])
        self.assertEqual(wb["A_Chi_tieu"].page_setup.orientation, "portrait")
        self.assertEqual(wb["D_Chua_hoan_tat"].page_setup.orientation, "landscape")
        for ws in wb.worksheets:
            self.assertEqual(int(ws.page_setup.paperSize), int(ws.PAPERSIZE_A4), ws.title)
            self.assertTrue(ws.print_area, ws.title)


class TongHopTest(CoSo):
    def test_phan_a_khong_cong_cheo_nhom(self):
        _, wb = self.xuat_ok("tong-hop")
        ws = wb["A_Chi_tieu"]
        h = next(r for r in range(1, 30) if ws.cell(r, 1).value == "Nhóm đối tượng")
        bang = [[c.value for c in hang] for hang in ws.iter_rows(min_row=h + 1, max_row=h + 3, max_col=10)]
        self.assertEqual(bang, [
            ["CBCS", 1, 6, 5, 2, 1, "—", 1, 6, 1],
            ["Khách", 0, 6, 4, 2, 1, 1, 0, 4, 0],
            ["Phương tiện", 0, 5, 4, 1, 0, 1, 0, 3, 0],
        ])
        self.assertEqual(ws.cell(h + 4, 1).value, "(Không có dòng tổng cộng giữa các nhóm.)")
        gia_tri = {c.value for hang in ws.iter_rows() for c in hang}
        self.assertFalse({"Tổng cộng", "Tổng"} & gia_tri)
        tieu_de = [c.value for c in ws[h] if c.value]
        self.assertTrue(all("lượt" in x or "đối tượng" in x or "bản ghi" in x for x in tieu_de[1:-1]), tieu_de)

    def test_phan_b_dong_cuoi_la_duy_nhat_ca_ky(self):
        _, wb = self.xuat_ok("tong-hop", tu="2026-10-24T00:00", den="2026-10-26T00:00")
        ws = wb["B_Theo_ngay"]
        h = next(r for r in range(1, 30) if ws.cell(r, 1).value == "Ngày")
        bang = [[c.value for c in hang] for hang in ws.iter_rows(min_row=h + 1, max_row=h + 3, max_col=7)]
        self.assertEqual(bang[0], [datetime.datetime(2026, 10, 24), 2, 2, 0, 0, 0, 0])
        self.assertEqual(bang[1], [datetime.datetime(2026, 10, 25), 6, 6, 6, 4, 5, 3])
        self.assertEqual(bang[2][0], "Cộng / Duy nhất cả kỳ (cột đối tượng không phải tổng các dòng)")
        self.assertEqual(bang[2][1:], [8, 6, 6, 4, 5, 3])

    def test_phan_d_an_dinh_danh_neu_khong_co_quyen_nhom(self):
        lanh_dao = tao_nguoi_dung("lanhdao01", quyen=["xem_so_tong_hop", "xuat_so_tong_hop", "xem_moi_cong"])
        _, wb = self.xuat_ok("tong-hop", nguoi=lanh_dao)
        ws = wb["D_Chua_hoan_tat"]
        noi_dung = " ".join(str(c.value) for hang in ws.iter_rows() for c in hang if c.value)
        for ten in ("Phạm Thị Dung", "Ông Quang", "Lý Thị Sen", "29X1-111.11"):
            self.assertNotIn(ten, noi_dung)
        self.assertEqual(noi_dung.count("(không có quyền xem chi tiết"), 5)


class AnToanDuLieuTest(CoSo):
    def test_cong_thuc_khong_duoc_thuc_thi(self):
        nguy_hiem = ["=HYPERLINK(\"http://x\",\"bấm\")", "+1+1", "-2+3", "@SUM(A1)", "=1+1", "\t=1",
                     "＝cmd|' /C calc'!A0"]
        k = Khach.objects.create(ho_ten=nguy_hiem[0], co_quan=nguy_hiem[1], so_giay_to="", loai_giay_to="")
        luot = Luot.objects.create(
            nhom="KHACH", chieu="VR", khach=k, mo_thoi_diem=t("12:00 25/10"), mo_cong=self.ds.c1,
            mo_nguoi_ghi=self.ds.tb01, mo_ghi_nhan_luc=t("12:00 25/10"), muc_dich=nguy_hiem[2], nguoi_tiep=nguy_hiem[3], ma_doan=nguy_hiem[4],
            so_the=nguy_hiem[5], ghi_chu=nguy_hiem[6],
        )
        phan_hoi, wb = self.xuat_ok("khach")
        dong = next(h for h in wb["SO02"].iter_rows() if h[-1].value == luot.pk)
        o_theo_gia_tri = {o.value: o for o in dong if isinstance(o.value, str)}
        for chuoi in nguy_hiem[:6]:
            o = o_theo_gia_tri[chuoi]  # Giữ nguyên nội dung, không thêm/bớt ký tự.
            self.assertEqual(o.data_type, "s", chuoi)
            self.assertTrue(o.quotePrefix, chuoi)
        self.assertEqual(o_theo_gia_tri["＝cmd|' /C calc'!A0"].data_type, "s")
        for ten, xml in xml_trong_tep(phan_hoi.content).items():
            self.assertNotIn("<f>", xml, ten)
            self.assertNotIn("<f ", xml, ten)

    def test_giu_so_0_dau_ma(self):
        don_vi = DonVi.objects.first()
        cb = CanBo.objects.create(ma="000123", ho_ten="Nguyễn Văn Không", don_vi=don_vi)
        Luot.objects.create(nhom="CBCS", chieu="RV", can_bo=cb, mo_thoi_diem=t("12:00 25/10"), mo_cong=self.ds.c1,
                            mo_nguoi_ghi=self.ds.tb01)
        _, wb = self.xuat_ok()
        ws = wb["SO01"]
        o = next(h[1] for h in ws.iter_rows() if h[1].value == "000123")
        self.assertEqual((o.data_type, o.number_format), ("s", "@"))
        _, wb = self.xuat_ok("khach", nguoi=self.ds.chihuy, du_lieu_day_du="on")
        so = [d["Số giấy tờ"] for d in doc_bang(wb["SO02"])[1]]
        self.assertIn("001090000111", so)
        self.assertIn("001 090 000 111", so)

    def test_khong_cat_du_lieu_dai(self):
        dai = "Nội dung rất dài có dấu tiếng Việt. " * 55  # khoảng 2000 ký tự
        Luot.objects.filter(pk=self.ds.l["C1"].pk).update(ly_do_chi_tiet=dai[:500], ghi_chu=dai.strip())
        _, wb = self.xuat_ok()
        c1 = next(d for d in doc_bang(wb["SO01"])[1] if d["Mã lượt"] == self.ds.l["C1"].pk)
        self.assertEqual(c1["Ghi chú"], dai.strip())
        self.assertEqual(c1["Nội dung"], dai[:500])
        o = next(h for h in wb["SO01"].iter_rows() if h[-1].value == self.ds.l["C1"].pk)[-2]
        self.assertTrue(o.alignment.wrap_text)

    def test_vuot_gioi_han_o_thi_that_bai_khong_cat(self):
        Luot.objects.filter(pk=self.ds.l["C1"].pk).update(ghi_chu="x" * 40000)
        phan_hoi = self.xuat()
        self.assertEqual(phan_hoi.status_code, 500)
        self.assertIn("vượt giới hạn 32767 ký tự", phan_hoi.content.decode())
        self.assertEqual(LichSuXuat.objects.get().ket_qua, LichSuXuat.KetQua.THAT_BAI)

    def test_ky_tu_dieu_khien_khong_lam_hong_tep(self):
        Luot.objects.filter(pk=self.ds.l["C1"].pk).update(ghi_chu="a\x01b\x1fc")
        _, wb = self.xuat_ok()
        c1 = next(d for d in doc_bang(wb["SO01"])[1] if d["Mã lượt"] == self.ds.l["C1"].pk)
        self.assertEqual(c1["Ghi chú"], "a�b�c")

    def test_tieng_viet(self):
        don_vi = DonVi.objects.first()
        ten = ["Nguyễn Thị Ánh Tuyết", "Ưng Hoàng Phúc", "Ông Ích Khiêm", "Bùi Thị Ơn"]
        for i, x in enumerate(ten):
            cb = CanBo.objects.create(ma=f"VN{i}", ho_ten=unicodedata.normalize("NFC", x), don_vi=don_vi)
            Luot.objects.create(nhom="CBCS", chieu="RV", can_bo=cb, mo_thoi_diem=t("12:00 25/10"),
                                mo_cong=self.ds.c1, mo_nguoi_ghi=self.ds.tb01, ly_do_chi_tiet="Đi công tác Đà Lạt")
        _, wb = self.xuat_ok(sap_xep="doi_tuong")
        dong = doc_bang(wb["SO01"])[1]
        theo_ten = [d["Họ và tên"] for d in dong]
        for x in ten:
            self.assertIn(x, theo_ten)
        self.assertEqual(wb["SO01"].cell(3, 1).value, "SỔ THEO DÕI CBCS RA/VÀO")
        # Sắp theo tên: An, An, Bình, Cường, Dung, Đức, Hưng, Khiêm, Ơn, Phúc, Tuyết
        self.assertEqual([x.split()[-1] for x in theo_ten],
                         ["An", "An", "Bình", "Cường", "Dung", "Đức", "Hưng", "Khiêm", "Ơn", "Phúc", "Tuyết"])

    def test_du_lieu_ca_nhan(self):
        phan_hoi, wb = self.xuat_ok("khach")
        so = [d["Số giấy tờ"] for d in doc_bang(wb["SO02"])[1]]
        self.assertEqual(so[0], "********0111")
        self.assertEqual(so[3], "********0111")  # "001 090 000 111": che trên số đã chuẩn hoá.
        toan_bo = " ".join(xml_trong_tep(phan_hoi.content).values())
        for bi_mat in ("001090000111", "0912345678", "001190000444"):
            self.assertNotIn(bi_mat, toan_bo)
        self.assertNotIn("Số điện thoại", toan_bo)
        self.assertIn("Dữ liệu cá nhân: số giấy tờ đã che", [wb["SO02"].cell(h, 1).value for h in range(1, 12)])

    def test_phu_luc_dinh_chinh_che_so_giay_to(self):
        self.ds.l["K6"].dinh_chinh("khach.so_giay_to", "001190000445", "Sai một chữ số", self.ds.tb01)
        phan_hoi, wb = self.xuat_ok("khach")
        _, dc = doc_bang(wb["Dinh_chinh"], 4)
        self.assertEqual((dc[0]["Giá trị cũ"], dc[0]["Giá trị mới"]), ("********0444", "********0445"))
        toan_bo = " ".join(xml_trong_tep(phan_hoi.content).values())
        self.assertNotIn("001190000444", toan_bo)
        self.assertNotIn("001190000445", toan_bo)


class PhanQuyenTest(CoSo):
    def test_chua_dang_nhap(self):
        phan_hoi = self.client.get(url("cbcs"), S1)
        self.assertEqual(phan_hoi.status_code, 302)
        self.assertFalse(LichSuXuat.objects.exists())

    def test_khong_co_quyen_xem_hoac_xuat(self):
        chi_xem = tao_nguoi_dung("xem01", quyen=["xem_so_cbcs", "xem_moi_cong"])
        khong_gi = tao_nguoi_dung("khonggi")
        for nguoi, so in [(chi_xem, "cbcs"), (khong_gi, "cbcs"), (chi_xem, "khach"), (khong_gi, "tong-hop")]:
            with self.subTest(nguoi=nguoi.username, so=so):
                self.assertEqual(self.xuat(so, nguoi).status_code, 403)
        # Bị từ chối kể cả khi tham số sai — không lộ việc kiểm tra tham số.
        self.assertEqual(self.xuat("cbcs", khong_gi, den="khong-hop-le").status_code, 403)
        self.assertFalse(LichSuXuat.objects.exists())

    def test_tuy_chon_can_quyen_rieng(self):
        self.assertEqual(self.xuat(hien_huy="on").status_code, 403)
        self.assertEqual(self.xuat("khach", du_lieu_day_du="on").status_code, 403)
        self.assertFalse(LichSuXuat.objects.exists())
        _, wb = self.xuat_ok(nguoi=self.ds.chihuy, hien_huy="on")
        dong = doc_bang(wb["SO01"])[1]
        c7 = next(d for d in dong if d["Mã lượt"] == self.ds.l["C7"].pk)
        self.assertEqual((c7["STT"], c7["Trạng thái"], c7["Ghi chú"]), ("—", "ĐÃ HUỶ", "Lý do huỷ: Nhập nhầm người"))

    def test_gioi_han_theo_cong(self):
        tb = tao_nguoi_dung("cong2", quyen=["xem_so_cbcs", "xuat_so_cbcs"])
        phan_cong(tb, self.ds.c2)
        _, wb = self.xuat_ok(nguoi=tb)
        self.assertEqual([d["Mã lượt"] for d in doc_bang(wb["SO01"])[1]],
                         [self.ds.l["C2"].pk, self.ds.l["C4"].pk])

    def test_so_khong_ton_tai(self):
        self.client.force_login(self.ds.tb01)
        self.assertEqual(self.client.get("/so/khong-co/xuat-excel/", S1).status_code, 404)

    def test_trang_danh_sach_chi_hien_so_duoc_xem(self):
        chi_xem = tao_nguoi_dung("xemcbcs", quyen=["xem_so_cbcs"])
        self.client.force_login(chi_xem)
        noi_dung = self.client.get(reverse("baocao:danh_sach")).content.decode()
        self.assertIn("SO-01", noi_dung)
        self.assertNotIn("SO-02", noi_dung)
        self.assertIn("Bạn không có quyền xuất sổ này.", noi_dung)


class KhoangThoiGianTest(CoSo):
    def test_khoang_khong_hop_le(self):
        truong_hop = [
            {"tu": "2026-10-26T00:00", "den": "2026-10-25T00:00"},
            {"tu": "2026-10-25T00:00", "den": "2026-10-25T00:00"},
            {"tu": "2025-01-01T00:00", "den": "2026-01-03T00:00"},
            {"tu": "2026-10-25T00:00", "den": ""},
            {"tu": "25-10-2026 abc", "den": "2026-10-26T00:00"},
            {"tu": "", "den": "", "ky": "khong_co"},
        ]
        for tham_so in truong_hop:
            with self.subTest(tham_so):
                phan_hoi = self.xuat(**tham_so)
                self.assertEqual(phan_hoi.status_code, 400)
                self.assertIn("Tham số không hợp lệ", phan_hoi.content.decode())
        self.assertFalse(LichSuXuat.objects.exists())

    def test_lua_chon_nhanh(self):
        phan_hoi, wb = self.xuat_ok(tu="", den="", ky="hom_qua")  # đồng hồ 27/10 → kỳ 26/10
        self.assertEqual(wb["SO01"].cell(4, 1).value, "Từ 00:00 ngày 26/10/2026 đến hết ngày 26/10/2026")

    def test_ky_chua_ket_thuc(self):
        with mock.patch("baocao.dong_ho.hien_tai", return_value=t("12:00 25/10")):
            _, wb = self.xuat_ok()
        self.assertEqual(wb["SO01"].cell(6, 1).value,
                         "Số liệu tính đến: 12:00:00 ngày 25/10/2026 (kỳ chưa kết thúc)")


class LichSuXuatTest(CoSo):
    def test_ghi_lich_su_thanh_cong(self):
        phan_hoi, wb = self.xuat_ok(cong=self.ds.c1.pk)
        ls = LichSuXuat.objects.get()
        self.assertEqual((ls.ma_so, ls.tu, ls.den, ls.nguoi_xuat, ls.ket_qua),
                         ("SO-01", t("00:00 25/10"), t("00:00 26/10"), self.ds.tb01, "THANH_CONG"))
        self.assertEqual((ls.bat_dau_luc, ls.thoi_diem_trich_xuat, ls.thoi_diem_chot),
                         (GIO_S1, GIO_S1, t("00:00 26/10")))
        self.assertEqual(ls.sha256, hashlib.sha256(phan_hoi.content).hexdigest())
        self.assertEqual(ls.ten_file, phan_hoi["Content-Disposition"].split('"')[1])
        self.assertEqual(ls.dieu_kien_loc["cong"], ["C1"])
        self.assertEqual(ls.nguoi_lap_so, "Nguyễn Văn X")
        self.assertEqual(ls.chi_tieu["so_luot"], 6)

    def test_ma_trich_xuat_tang_dan_khong_trung(self):
        ma = [self.xuat_ok()[0]["X-Ma-Trich-Xuat"] for _ in range(3)]
        self.assertEqual(ma, ["SO01-20261027-0001", "SO01-20261027-0002", "SO01-20261027-0003"])
        self.assertEqual(self.xuat_ok("khach")[0]["X-Ma-Trich-Xuat"], "SO02-20261027-0004")

    def test_khong_ghi_thanh_cong_khi_luu_tep_that_bai(self):
        with mock.patch("baocao.excel.Workbook.save", side_effect=OSError("đĩa đầy")):
            phan_hoi = self.xuat()
        self.assertEqual(phan_hoi.status_code, 500)
        self.assertNotEqual(phan_hoi.get("Content-Type"), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        ls = LichSuXuat.objects.get()
        self.assertEqual(ls.ket_qua, LichSuXuat.KetQua.THAT_BAI)
        self.assertIn("OSError: đĩa đầy", ls.loi)
        self.assertEqual((ls.sha256, ls.ten_file, ls.so_dong, ls.thoi_diem_trich_xuat), ("", "", None, None))
        self.assertIsNotNone(ls.ket_thuc_luc)
        self.assertIn(ls.ma_trich_xuat, phan_hoi.content.decode())

    def test_khong_ghi_thanh_cong_khi_lap_so_lieu_that_bai(self):
        with mock.patch("baocao.so_lieu.tinh_chi_tieu", side_effect=RuntimeError("lỗi giả lập")):
            phan_hoi = self.xuat("tong-hop")
        self.assertEqual(phan_hoi.status_code, 500)
        self.assertEqual(LichSuXuat.objects.get().ket_qua, LichSuXuat.KetQua.THAT_BAI)
        # Lần sau thành công cấp mã mới, không dùng lại mã đã thất bại.
        self.assertEqual(self.xuat_ok("tong-hop")[0]["X-Ma-Trich-Xuat"], "SO04-20261027-0002")


class SoTrongVaLonTest(CoSo):
    def test_ky_khong_co_du_lieu(self):
        for so, sheet in (("cbcs", "SO01"), ("khach", "SO02"), ("phuong-tien", "SO03")):
            with self.subTest(so):
                _, wb = self.xuat_ok(so, tu="2026-01-01T00:00", den="2026-01-02T00:00")
                ws = wb[sheet]
                h = hang_tieu_de(ws)
                self.assertEqual(ws.cell(h + 1, 1).value, "Không có lượt nào trong kỳ theo điều kiện đã chọn.")
                tong = {ws.cell(r, 2).value: ws.cell(r, 6).value for r in range(h, ws.max_row + 1)}
                self.assertEqual(tong["Số lượt trong sổ"], 0)
                self.assertIn("NGƯỜI LẬP SỔ", {ws.cell(r, 2).value for r in range(h, ws.max_row + 1)})
        _, wb = self.xuat_ok("tong-hop", tu="2026-01-01T00:00", den="2026-01-02T00:00")
        self.assertEqual(wb["D_Chua_hoan_tat"].cell(hang_d(wb) + 1, 1).value, "Không có lượt chưa hoàn tất.")
        self.assertEqual(list(LichSuXuat.objects.values_list("so_dong", flat=True).order_by("bat_dau_luc", "id")),
                         [0, 0, 0, 3 + 1 + 0 + 0 + 3])

    def test_du_lieu_lon_va_doi_chieu_so_dong(self):
        don_vi = DonVi.objects.first()
        can_bo = CanBo.objects.bulk_create(
            [CanBo(ma=f"{i:05d}", ho_ten=f"Cán bộ Thử {i}", don_vi=don_vi) for i in range(200)]
        )
        bat_dau = t("00:00 01/10")
        luot = []
        for i in range(10_000):
            mo = bat_dau + datetime.timedelta(minutes=4 * i)  # trải đều tháng 10, có lượt qua nửa đêm
            dong = None if i % 100 == 0 else mo + datetime.timedelta(minutes=90)
            luot.append(Luot(
                nhom="CBCS", chieu="RV", can_bo=can_bo[i % 200], mo_thoi_diem=mo, mo_cong=self.ds.c1,
                mo_ghi_nhan_luc=mo, mo_nguoi_ghi=self.ds.tb01, dong_thoi_diem=dong,
                dong_cong=self.ds.c1 if dong else None, dong_ghi_nhan_luc=dong,
                dong_nguoi_ghi=self.ds.tb01 if dong else None, ly_do_chi_tiet=f"Lượt thử nghiệm số {i}",
            ))
        Luot.objects.bulk_create(luot, batch_size=2000)
        DinhChinh.objects.all().delete()

        ky = {"tu": "2026-10-01T00:00", "den": "2026-11-01T00:00"}
        bat_dau_do = time.monotonic()
        with mock.patch("baocao.dong_ho.hien_tai", return_value=t("09:00 02/11")):  # kỳ đã kết thúc
            phan_hoi, wb = self.xuat_ok(**ky)
        thoi_gian = time.monotonic() - bat_dau_do
        ws = wb["SO01"]
        _, dong = doc_bang(ws)

        # Đối chiếu với truy vấn độc lập theo BR-KY-05.
        chot = t("00:00 01/11")
        from django.db.models import Q
        mong_doi = Luot.objects.filter(
            Q(dong_thoi_diem__isnull=True) | Q(dong_thoi_diem__gte=t("00:00 01/10")),
            nhom="CBCS", da_huy=False, mo_thoi_diem__lt=chot,
        ).count()
        ls = LichSuXuat.objects.get()
        self.assertEqual(len(dong), mong_doi)
        self.assertEqual(mong_doi, 10_008)  # 10.000 lượt sinh thêm + 8 lượt hiệu lực của DS-01
        self.assertEqual(ls.so_dong, mong_doi)
        self.assertEqual(ls.chi_tieu["so_luot"], mong_doi)
        self.assertEqual([d["STT"] for d in dong], list(range(1, mong_doi + 1)))
        self.assertEqual(len({d["Mã lượt"] for d in dong}), mong_doi)
        self.assertEqual(ws.auto_filter.ref, f"A{hang_tieu_de(ws)}:{ws.auto_filter.ref.split(':')[1][:2]}"
                         f"{hang_tieu_de(ws) + mong_doi}")
        self.assertLess(thoi_gian, 60, f"Xuất 10.000 lượt mất {thoi_gian:.1f} giây")
        print(f"\n[du_lieu_lon] {mong_doi} dòng, {len(phan_hoi.content) / 1024:.0f} KiB, {thoi_gian:.1f} giây")


def hang_d(wb):
    ws = wb["D_Chua_hoan_tat"]
    return next(r for r in range(1, 30) if ws.cell(r, 1).value == "Nhóm")


class DocNhatQuanTest(TransactionTestCase):
    """Chạy ngoài giao dịch kiểm thử để đi qua nhánh SET TRANSACTION REPEATABLE READ."""

    def test_xuat_trong_giao_dich_doc_nhat_quan(self):
        ds = DS01()
        self.client.force_login(ds.tb01)
        from django.db import connection

        cau_lenh = []

        def ghi_lai(execute, sql, params, many, context):
            cau_lenh.append(sql)
            return execute(sql, params, many, context)

        with connection.execute_wrapper(ghi_lai):
            phan_hoi = self.client.get(url("tong-hop"), S1)
        self.assertEqual(phan_hoi.status_code, 200)
        self.assertIn("SET TRANSACTION ISOLATION LEVEL REPEATABLE READ READ ONLY", cau_lenh)
        self.assertEqual(LichSuXuat.objects.get().ket_qua, "THANH_CONG")

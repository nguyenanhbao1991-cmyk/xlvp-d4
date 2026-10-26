"""Màn hình "Xuất sổ theo dõi": từ giao diện tới cơ sở dữ liệu và tệp đầu ra."""
import datetime
import hashlib
import io
import re
import uuid
from unittest import mock

from django.contrib.messages import get_messages
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from docx import Document

from ravao.models import CanBo, Luot

from ..ky import KyBaoCao
from ..models import BanChotXuat, CauHinhSo, LichSuXuat
from .du_lieu import DS01, doc_bang, mo_workbook, phan_cong, t, tao_nguoi_dung
from .test_xuat_word_pdf import bang_word, gon, trang_pdf

URL = "/so/xuat/"
GIO = t("09:00 27/10")
UTC = datetime.timezone.utc


class CoSo(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.ds = DS01()
        CauHinhSo.objects.create(co_quan_chu_quan="Cơ quan thử nghiệm", ten_don_vi="Đơn vị thử nghiệm")

    def setUp(self):
        dong_ho = mock.patch("baocao.dong_ho.hien_tai", return_value=GIO)
        dong_ho.start()
        self.addCleanup(dong_ho.stop)

    def du_lieu(self, so="SO-01", dinh_dang="xlsx", che_do="ngay", ngay="2026-10-25", loc=None, **them):
        du_lieu = {"loai_so": so, "dinh_dang": dinh_dang, "che_do_ky": che_do, "ngay": ngay, **them}
        tien_to = so.replace("-", "").lower()
        for k, v in (loc or {}).items():
            du_lieu[f"{tien_to}-{k}"] = v
        return du_lieu

    def post(self, hanh_dong, nguoi=None, **kw):
        self.client.force_login(nguoi or self.ds.tb01)
        return self.client.post(URL, {**self.du_lieu(**kw), "hanh_dong": hanh_dong})

    def xuat_va_tai(self, nguoi=None, **kw):
        """Bấm “Xuất tệp”, theo chuyển hướng, rồi tải tệp bằng liên kết trên màn hình."""
        phan_hoi = self.post("xuat", nguoi, **kw)
        self.assertEqual(phan_hoi.status_code, 302, self.thong_bao(phan_hoi))
        ma_tai = phan_hoi["Location"].split("ket_qua=")[1]
        trang = self.client.get(phan_hoi["Location"])
        lien_ket = reverse("baocao:tai_tep", args=[ma_tai])
        self.assertContains(trang, lien_ket)
        tep = self.client.get(lien_ket)
        self.assertEqual(tep.status_code, 200)
        ls = LichSuXuat.objects.get(ma_tai=ma_tai)
        return ls, tep.content, trang

    @staticmethod
    def thong_bao(phan_hoi):
        return [str(m) for m in get_messages(phan_hoi.wsgi_request)]

    def ma(self, *ten):
        return [self.ds.l[x].pk for x in ten]

    @staticmethod
    def ma_trong_html(noi_dung):
        """Mã lượt in dưới STT trong bảng xem trước (ô "1<br>#84")."""
        return [int(x) for x in re.findall(r"<td>\d+<br>#(\d+)</td>", noi_dung.decode())]


# ---------------------------------------------------------------------------
# Giao diện
# ---------------------------------------------------------------------------


class GiaoDienTest(CoSo):
    def test_trang_chi_hien_so_duoc_xem(self):
        self.client.force_login(tao_nguoi_dung("xemcbcs", quyen=["xem_so_cbcs"]))
        noi_dung = self.client.get(URL).content.decode()
        self.assertIn('value="SO-01"', noi_dung)
        self.assertNotIn('value="SO-02"', noi_dung)
        self.assertIn("Bạn chỉ có quyền xem trước sổ này, không có quyền xuất.", noi_dung)
        self.assertIn('value="2026-10-27"', noi_dung)  # mặc định hôm nay theo giờ Việt Nam
        self.client.force_login(tao_nguoi_dung("khonggi"))
        self.assertEqual(self.client.get(URL).status_code, 403)
        self.client.logout()
        self.assertEqual(self.client.get(URL).status_code, 302)

    def test_khong_phu_thuoc_internet(self):
        self.client.force_login(self.ds.tb01)
        noi_dung = self.client.get(URL).content.decode()
        self.assertNotRegex(noi_dung, r"(src|href)=\"https?://")


# ---------------------------------------------------------------------------
# Xem trước
# ---------------------------------------------------------------------------


class XemTruocTest(CoSo):
    def test_xem_truoc_so_chi_tiet(self):
        phan_hoi = self.post("xem_truoc")
        self.assertEqual(phan_hoi.status_code, 200)
        noi_dung = phan_hoi.content.decode()
        self.assertIn("XEM TRƯỚC — CHƯA PHẢI DỮ LIỆU ĐÃ CHỐT", noi_dung)
        self.assertIn('Số dòng sẽ xuất: <span id="so-dong-xem-truoc">7</span>', noi_dung)
        self.assertEqual(self.ma_trong_html(phan_hoi.content), self.ma("C3", "C8", "C1", "C6", "C2", "C5", "C4"))
        self.assertIn('name="so_dong_xem_truoc" value="7"', noi_dung)
        # Qua nửa đêm, chưa hoàn tất, chỉ tiêu.
        self.assertIn("23:30 24/10<br>Cổng 1<br>(từ kỳ trước)", noi_dung)
        self.assertIn("00:20 26/10<br>Cổng 1<br>(sau kỳ)", noi_dung)
        self.assertRegex(noi_dung, r'<tr class="chua-hoan-tat">.*?#%d.*?≥ 08:00.*?Chưa về – Quá hạn'
                         % self.ds.l["C5"].pk)
        self.assertRegex(noi_dung, r"Chưa hoàn tất tại thời điểm chốt</td><td class=\"so\">2</td>")
        # Xem trước không cấp mã, không ghi lịch sử, không lưu bản chốt.
        self.assertFalse(LichSuXuat.objects.exists())
        self.assertFalse(BanChotXuat.objects.exists())

    def test_xem_truoc_so_tong_hop(self):
        phan_hoi = self.post("xem_truoc", so="SO-04")
        noi_dung = phan_hoi.content.decode()
        gon_html = re.sub(r"\s*\n\s*", "", noi_dung)
        for nhom, so in (("CBCS", [1, 6, 5, 2, 1, 6]), ("Khách", [0, 6, 4, 2, 1, 4]), ("Phương tiện", [0, 5, 4, 1, 0, 3])):
            self.assertIn(f"<td>{nhom}</td>" + "".join(f'<td class="so">{x}</td>' for x in so), gon_html)
        # A 3 nhóm + B 1 ngày + C 2 cổng + D 5 lượt chưa hoàn tất + E 3 nhóm.
        self.assertIn('Số dòng sẽ xuất: <span id="so-dong-xem-truoc">14</span>', noi_dung)
        self.assertFalse(LichSuXuat.objects.exists())

    def test_khong_co_du_lieu_khong_tao_so_lieu_gia(self):
        phan_hoi = self.post("xem_truoc", ngay="2026-01-01")
        noi_dung = phan_hoi.content.decode()
        self.assertIn('Số dòng sẽ xuất: <span id="so-dong-xem-truoc">0</span>', noi_dung)
        self.assertIn("Không có lượt nào trong kỳ theo điều kiện đã chọn.", noi_dung)
        self.assertEqual(self.ma_trong_html(phan_hoi.content), [])
        ls, tep, _ = self.xuat_va_tai(ngay="2026-01-01")
        self.assertEqual(ls.so_dong, 0)
        self.assertEqual(doc_bang(mo_workbook(tep)["SO01"])[1], [])
        self.assertEqual(ls.ban_chot.du_lieu["dong"], [])
        self.assertEqual(ls.ban_chot.du_lieu["chi_tieu"]["so_luot"], 0)

    def test_canh_bao_gioi_han_ban_in(self):
        phan_hoi = self.post("xem_truoc", dinh_dang="pdf", che_do="khoang", tu_ngay="2026-09-01", den_ngay="2026-10-25")
        self.assertContains(phan_hoi, "Kỳ dài hơn 31 ngày: bản PDF sẽ bị từ chối, chỉ xuất được Excel.")


# ---------------------------------------------------------------------------
# Xuất, bản chốt, đối chiếu
# ---------------------------------------------------------------------------


class XuatVaDoiChieuTest(CoSo):
    THU_TU = ("C3", "C8", "C1", "C6", "C2", "C5", "C4")

    def test_xuat_ba_dinh_dang_doi_chieu_so_dong(self):
        xem = self.post("xem_truoc")
        ma_xem = self.ma_trong_html(xem.content)
        for dinh_dang in ("xlsx", "docx", "pdf"):
            with self.subTest(dinh_dang):
                ls, tep, trang = self.xuat_va_tai(dinh_dang=dinh_dang)
                self.assertEqual(ls.sha256, hashlib.sha256(tep).hexdigest())
                self.assertEqual((ls.ket_qua, ls.dinh_dang, ls.phien_ban_mau, ls.so_dong),
                                 ("THANH_CONG", dinh_dang, "1.0", 7))
                ban_chot = [d["ma_luot"] for d in ls.ban_chot.du_lieu["dong"]]
                if dinh_dang == "xlsx":
                    wb = mo_workbook(tep)
                    trong_tep = [d["Mã lượt"] for d in doc_bang(wb["SO01"])[1]]
                    self.assertIn("Mẫu nội bộ — phiên bản 1.0", [o.value for o in wb["SO01"][2]])
                elif dinh_dang == "docx":
                    trong_tep = [int(h.cells[0].text.split("#")[1]) for h in bang_word(Document(io.BytesIO(tep))).rows[1:]]
                else:
                    # Bảng đứng trước phụ lục đính chính: lấy lần xuất hiện đầu tiên của mỗi mã.
                    trong_tep = list(dict.fromkeys(int(x) for x in re.findall(r"#(\d+)", "\n".join(trang_pdf(tep)[1]))))
                self.assertEqual(trong_tep, ma_xem)
                self.assertEqual(ban_chot, ma_xem)
                self.assertEqual(ma_xem, self.ma(*self.THU_TU))
                noi_dung = trang.content.decode()
                self.assertIn(f"Đã tạo tệp {ls.ten_file} — mã trích xuất {ls.ma_trich_xuat}, 7 dòng dữ liệu", noi_dung)
                self.assertIn("Kết quả: Thành công", noi_dung)
                # Thời điểm trong thông báo phải theo giờ Việt Nam, không phải UTC lưu trong CSDL.
                self.assertIn("số liệu tính đến 00:00:00 26/10/2026 (giờ Việt Nam)", noi_dung)
                self.assertNotIn("17:00:00 25/10/2026", noi_dung)

    def test_ban_chot_luu_so_lieu_va_qua_nua_dem(self):
        ls, _, _ = self.xuat_va_tai()
        du_lieu = ls.ban_chot.du_lieu
        self.assertEqual(du_lieu["chi_tieu"], {k: v for k, v in ls.chi_tieu.items()})
        dong = {d["ma_luot"]: d for d in du_lieu["dong"]}
        c3, c4, c5 = (dong[self.ds.l[x].pk] for x in ("C3", "C4", "C5"))
        self.assertEqual((c3["tu_ky_truoc"], c3["trang_thai"], c3["thoi_luong_phut"]), (True, "HT", 105))
        self.assertEqual((c4["sau_ky"], c4["trang_thai"], c4["thoi_luong_phut"]), (True, "HTS", 95))
        self.assertEqual((c5["trang_thai"], c5["qua_han"], c5["thoi_luong_phut"], c5["thoi_luong_toi_thieu"]),
                         ("MO", True, 480, True))
        self.assertEqual(du_lieu["phien_ban_mau"], "1.0")
        # Bản chốt không chứa họ tên hay số giấy tờ.
        self.assertNotIn("Phạm Thị Dung", str(du_lieu))

    def test_ky_theo_khoang_ngay_va_mui_gio(self):
        ls, _, _ = self.xuat_va_tai(che_do="khoang", tu_ngay="2026-10-24", den_ngay="2026-10-25")
        self.assertEqual(ls.tu.astimezone(UTC), datetime.datetime(2026, 10, 23, 17, 0, tzinfo=UTC))
        self.assertEqual(ls.den.astimezone(UTC), datetime.datetime(2026, 10, 25, 17, 0, tzinfo=UTC))
        ct = ls.ban_chot.du_lieu["chi_tieu"]
        self.assertEqual((ct["tu_ky_truoc"], ct["phat_sinh"], ct["hoan_tat"], ct["chua_hoan_tat"]), (0, 8, 6, 2))
        # Cùng ngày bắt đầu và kết thúc là hợp lệ.
        ls, _, _ = self.xuat_va_tai(che_do="khoang", tu_ngay="2026-10-25", den_ngay="2026-10-25")
        self.assertEqual(ls.so_dong, 7)

    def test_ky_theo_thang(self):
        ls, tep, _ = self.xuat_va_tai(che_do="thang", thang="2026-10", ngay="")
        self.assertEqual((timezone.localtime(ls.tu).date(), timezone.localtime(ls.den).date()),
                         (datetime.date(2026, 10, 1), datetime.date(2026, 11, 1)))
        self.assertEqual(wb_dong := len(doc_bang(mo_workbook(tep)["SO01"])[1]), 8)  # C0…C8 trừ C7 huỷ
        self.assertEqual(ls.so_dong, wb_dong)
        self.assertIn("Từ 00:00 ngày 01/10/2026 đến hết ngày 31/10/2026", mo_workbook(tep)["SO01"].cell(4, 1).value)

    def test_ranh_gioi_nua_dem_gio_viet_nam(self):
        cb = self.ds.cb["CB006"]
        # 00:30 ngày 25/10 giờ Việt Nam = 17:30 ngày 24/10 UTC → thuộc ngày 25/10.
        trong = Luot.objects.create(nhom="CBCS", chieu="RV", can_bo=cb, mo_cong=self.ds.c1, mo_nguoi_ghi=self.ds.tb01,
                                    mo_thoi_diem=datetime.datetime(2026, 10, 24, 17, 30, tzinfo=UTC),
                                    mo_ghi_nhan_luc=datetime.datetime(2026, 10, 24, 17, 30, tzinfo=UTC),
                                    dong_thoi_diem=datetime.datetime(2026, 10, 24, 18, 0, tzinfo=UTC),
                                    dong_cong=self.ds.c1, dong_nguoi_ghi=self.ds.tb01,
                                    dong_ghi_nhan_luc=datetime.datetime(2026, 10, 24, 18, 0, tzinfo=UTC))
        # 23:00–23:59 ngày 24/10 giờ Việt Nam → không thuộc ngày 25/10 dù là ngày 24/10 theo UTC… và ngược lại.
        ngoai = Luot.objects.create(nhom="CBCS", chieu="RV", can_bo=cb, mo_cong=self.ds.c1, mo_nguoi_ghi=self.ds.tb01,
                                    mo_thoi_diem=datetime.datetime(2026, 10, 24, 16, 0, tzinfo=UTC),
                                    mo_ghi_nhan_luc=datetime.datetime(2026, 10, 24, 16, 0, tzinfo=UTC),
                                    dong_thoi_diem=datetime.datetime(2026, 10, 24, 16, 59, tzinfo=UTC),
                                    dong_cong=self.ds.c1, dong_nguoi_ghi=self.ds.tb01,
                                    dong_ghi_nhan_luc=datetime.datetime(2026, 10, 24, 16, 59, tzinfo=UTC))
        ma = self.ma_trong_html(self.post("xem_truoc").content)
        self.assertIn(trong.pk, ma)
        self.assertNotIn(ngoai.pk, ma)
        self.assertIn(ngoai.pk, self.ma_trong_html(self.post("xem_truoc", ngay="2026-10-24").content))

    def test_ky_khong_doi_khi_phien_kich_hoat_mui_gio_khac(self):
        with timezone.override("UTC"):
            ky = KyBaoCao.theo_ngay(datetime.date(2026, 10, 25), datetime.date(2026, 10, 25))
        self.assertEqual(ky.tu.astimezone(UTC), datetime.datetime(2026, 10, 24, 17, 0, tzinfo=UTC))

    def test_bo_loc_tu_giao_dien_den_tep(self):
        truong_hop = [
            ("SO-01", {"trang_thai": "CHUA_HOAN_TAT"}, ("C5", "C4"), "SO01", "Trạng thái: Chưa hoàn tất"),
            ("SO-01", {"cong": str(self.ds.c2.pk)}, ("C2", "C4"), "SO01", "Điểm kiểm soát: Cổng 2"),
            ("SO-02", {"ma_doan": "ĐK-01"}, ("K3", "K4"), "SO02", "Đoàn: ĐK-01"),
            ("SO-03", {"nguon": "CO_QUAN"}, ("P1", "P2"), "SO03", "Nguồn: CO_QUAN"),
        ]
        for so, loc, mong_doi, sheet, mo_ta in truong_hop:
            with self.subTest(so=so, loc=loc):
                ma_xem = self.ma_trong_html(self.post("xem_truoc", so=so, loc=loc).content)
                ls, tep, _ = self.xuat_va_tai(so=so, loc=loc)
                ws = mo_workbook(tep)[sheet]
                self.assertEqual([d["Mã lượt"] for d in doc_bang(ws)[1]], self.ma(*mong_doi))
                self.assertEqual(ma_xem, self.ma(*mong_doi))
                self.assertIn(mo_ta, ws.cell(5, 1).value)
                self.assertEqual(ws.cell(9, 1).value, "Sổ đã lọc — không phải toàn bộ dữ liệu của kỳ.")
                self.assertEqual(ls.so_dong, len(mong_doi))

    def test_so_tong_hop(self):
        so_dong_xem = int(re.search(r'id="so-dong-xem-truoc">(\d+)<',
                                    self.post("xem_truoc", so="SO-04", che_do="khoang", tu_ngay="2026-10-24",
                                              den_ngay="2026-10-25").content.decode()).group(1))
        ls, tep, _ = self.xuat_va_tai(so="SO-04", dinh_dang="pdf", che_do="khoang", tu_ngay="2026-10-24",
                                      den_ngay="2026-10-25")
        du_lieu = ls.ban_chot.du_lieu
        self.assertEqual(du_lieu["phan_b_cuoi"], {"CBCS": [8, 6], "KHACH": [6, 4], "PT": [5, 3]})
        self.assertEqual([d["ma_luot"] for d in du_lieu["phan_d"]["CBCS"]], self.ma("C5", "C4"))
        self.assertEqual(du_lieu["phan_a"]["CBCS"]["phat_sinh"], 8)
        self.assertEqual(ls.so_dong, so_dong_xem)
        self.assertEqual(ls.so_dong, sum(ls.chi_tieu["so_dong_theo_phan"].values()))
        chu = gon("\n".join(trang_pdf(tep)[1]))
        self.assertIn(gon("Cộng / Duy nhất cả kỳ"), chu)
        self.assertNotIn("Tổngcộng", chu)


class BanChotVaXemTruocTest(CoSo):
    def test_tai_lai_tra_dung_ban_chot_du_du_lieu_goc_thay_doi(self):
        ls, tep, _ = self.xuat_va_tai()
        du_lieu_truoc = ls.ban_chot.du_lieu
        self.ds.l["C1"].dinh_chinh("mo_thoi_diem", t("08:05 25/10"), "Sửa giờ", self.ds.tb02)
        Luot.objects.create(nhom="CBCS", chieu="RV", can_bo=self.ds.cb["CB006"], mo_thoi_diem=t("12:00 25/10"),
                            mo_cong=self.ds.c1, mo_ghi_nhan_luc=t("12:00 25/10"), mo_nguoi_ghi=self.ds.tb01)
        self.client.force_login(self.ds.tb01)
        tai_lai = self.client.get(reverse("baocao:tai_tep", args=[ls.ma_tai])).content
        self.assertEqual(tai_lai, tep)
        ls.ban_chot.refresh_from_db()
        self.assertEqual(ls.ban_chot.du_lieu, du_lieu_truoc)
        moi, tep_moi, _ = self.xuat_va_tai()
        self.assertEqual(moi.so_dong, 8)
        self.assertNotEqual(moi.sha256, ls.sha256)
        self.assertEqual(LichSuXuat.objects.get(pk=ls.pk).so_lan_tai, 2)

    def test_xem_truoc_khong_phai_ban_chot(self):
        xem = self.post("xem_truoc")
        so_dong_xem = re.search(r'name="so_dong_xem_truoc" value="(\d+)"', xem.content.decode()).group(1)
        Luot.objects.create(nhom="CBCS", chieu="RV", can_bo=self.ds.cb["CB006"], mo_thoi_diem=t("12:00 25/10"),
                            mo_cong=self.ds.c1, mo_ghi_nhan_luc=t("12:00 25/10"), mo_nguoi_ghi=self.ds.tb01)
        phan_hoi = self.post("xuat", so_dong_xem_truoc=so_dong_xem)
        thong_bao = self.thong_bao(phan_hoi)
        self.assertIn("Dữ liệu đã thay đổi kể từ lúc xem trước (xem trước: 7 dòng, bản xuất: 8 dòng). "
                      "Bản xuất vừa tạo là bản chốt.", thong_bao)
        ls = LichSuXuat.objects.get()
        self.assertEqual((ls.so_dong, len(ls.ban_chot.du_lieu["dong"])), (8, 8))

    def test_tep_luu_bi_sua_thi_khong_tai(self):
        ls, _, _ = self.xuat_va_tai()
        BanChotXuat.objects.filter(pk=ls.ban_chot.pk).update(noi_dung=b"da bi thay doi")
        phan_hoi = self.client.get(reverse("baocao:tai_tep", args=[ls.ma_tai]))
        self.assertEqual(phan_hoi.status_code, 500)
        self.assertIn("không khớp mã kiểm tra", phan_hoi.content.decode())


# ---------------------------------------------------------------------------
# Kiểm tra dữ liệu nhập, thất bại
# ---------------------------------------------------------------------------


class KiemTraNhapVaThatBaiTest(CoSo):
    def test_ngay_bat_dau_sau_ngay_ket_thuc(self):
        for hanh_dong in ("xem_truoc", "xuat"):
            phan_hoi = self.post(hanh_dong, che_do="khoang", tu_ngay="2026-10-26", den_ngay="2026-10-25")
            self.assertEqual(phan_hoi.status_code, 400)
            self.assertContains(phan_hoi, "Ngày bắt đầu không được sau ngày kết thúc.", status_code=400)
        self.assertFalse(LichSuXuat.objects.exists())

    def test_thieu_hoac_sai_ngay(self):
        truong_hop = [
            dict(ngay=""), dict(che_do="khoang", tu_ngay="2026-10-25", den_ngay=""),
            dict(che_do="thang", thang="2026-13"), dict(che_do="thang", thang=""), dict(ngay="25-10-2026x"),
            dict(che_do="khoang", tu_ngay="2025-01-01", den_ngay="2026-10-25"),  # quá 366 ngày
            dict(dinh_dang="odt"), dict(so="SO-09"),
        ]
        for kw in truong_hop:
            with self.subTest(kw):
                self.assertEqual(self.post("xuat", **kw).status_code, 400)
        self.assertEqual(self.client.post(URL, {**self.du_lieu(), "hanh_dong": "xoa"}).status_code, 400)
        self.assertFalse(LichSuXuat.objects.exists())

    def test_ky_qua_dai_cho_ban_in(self):
        phan_hoi = self.post("xuat", dinh_dang="docx", che_do="khoang", tu_ngay="2026-09-01", den_ngay="2026-10-25")
        self.assertContains(phan_hoi, "tối đa 31 ngày", status_code=400)
        self.assertFalse(LichSuXuat.objects.exists())
        # Theo tháng (31 ngày) vẫn xuất được Word.
        ls, _, _ = self.xuat_va_tai(dinh_dang="docx", che_do="thang", thang="2026-10")
        self.assertEqual(ls.ket_qua, "THANH_CONG")

    def test_tao_tep_that_bai(self):
        with mock.patch("baocao.in_pdf.tao_pdf", side_effect=RuntimeError("hỏng")):
            phan_hoi = self.post("xuat", dinh_dang="pdf")
        self.assertEqual(phan_hoi.status_code, 302)
        thong_bao = self.thong_bao(phan_hoi)
        ls = LichSuXuat.objects.get()
        self.assertIn(f"Xuất thất bại (lỗi khi tạo tệp, mã {ls.ma_trich_xuat}): Không tạo được tệp PDF.", thong_bao)
        self.assertEqual(ls.ket_qua, "THAT_BAI")
        self.assertFalse(BanChotXuat.objects.exists())
        trang = self.client.get(phan_hoi["Location"]).content.decode()
        self.assertIn("Kết quả: Thất bại", trang)
        self.assertNotIn(reverse("baocao:tai_tep", args=[ls.ma_tai]), trang)
        self.assertEqual(self.client.get(reverse("baocao:tai_tep", args=[ls.ma_tai])).status_code, 404)

    def test_luu_ban_chot_that_bai_thi_khong_ghi_thanh_cong(self):
        with mock.patch("baocao.dich_vu.BanChotXuat.objects.create", side_effect=RuntimeError("đĩa đầy")):
            self.post("xuat")
        ls = LichSuXuat.objects.get()
        self.assertEqual((ls.ket_qua, ls.sha256), ("THAT_BAI", ""))
        self.assertFalse(BanChotXuat.objects.exists())


# ---------------------------------------------------------------------------
# Phân quyền và tải tệp
# ---------------------------------------------------------------------------


class PhanQuyenTest(CoSo):
    def test_khong_du_quyen_thi_khong_tao_tep(self):
        chi_xem = tao_nguoi_dung("chixem", quyen=["xem_so_cbcs", "xem_moi_cong"])
        self.assertEqual(self.post("xem_truoc", nguoi=chi_xem).status_code, 200)
        phan_hoi = self.post("xuat", nguoi=chi_xem)
        self.assertEqual(phan_hoi.status_code, 403)
        self.assertContains(phan_hoi, "Không có quyền xuất SO-01", status_code=403)
        self.assertContains(phan_hoi, "Không có tệp nào được tạo.", status_code=403)
        # Sổ không được xem không có trong lựa chọn.
        self.assertEqual(self.post("xuat", nguoi=chi_xem, so="SO-02").status_code, 400)
        # Tuỳ chọn cần quyền riêng.
        self.assertEqual(self.post("xuat", loc={"hien_huy": "on"}).status_code, 403)
        self.assertEqual(self.post("xem_truoc", loc={"hien_huy": "on"}).status_code, 403)
        self.assertEqual(self.post("xuat", so="SO-02", loc={"du_lieu_day_du": "on"}).status_code, 403)
        self.assertFalse(LichSuXuat.objects.exists())
        self.assertFalse(BanChotXuat.objects.exists())

    def test_gioi_han_theo_cong(self):
        tb = tao_nguoi_dung("cong2", quyen=["xem_so_cbcs", "xuat_so_cbcs"])
        phan_cong(tb, self.ds.c2)
        self.assertEqual(self.ma_trong_html(self.post("xem_truoc", nguoi=tb).content), self.ma("C2", "C4"))
        ls, tep, _ = self.xuat_va_tai(nguoi=tb)
        self.assertEqual([d["Mã lượt"] for d in doc_bang(mo_workbook(tep)["SO01"])[1]], self.ma("C2", "C4"))

    def test_khong_tai_duoc_tep_cua_nguoi_khac(self):
        ls, _, _ = self.xuat_va_tai()
        khac = tao_nguoi_dung("trucban09", quyen=["xem_so_cbcs", "xuat_so_cbcs", "xem_moi_cong"])
        self.client.force_login(khac)
        duong_dan = reverse("baocao:tai_tep", args=[ls.ma_tai])
        self.assertEqual(self.client.get(duong_dan).status_code, 404)
        # Đoán đường dẫn: mã ngẫu nhiên, mã trích xuất, khoá chính, mã không đúng dạng.
        for doan in (f"/so/tai/{uuid.uuid4()}/", f"/so/tai/{ls.ma_trich_xuat}/", f"/so/tai/{ls.pk}/",
                     "/so/tai/../../admin/", f"/so/tai/{str(ls.ma_tai).upper()}x/"):
            with self.subTest(doan):
                self.assertEqual(self.client.get(doan).status_code, 404)
        # Không lộ kết quả của người khác qua tham số ket_qua.
        trang = self.client.get(f"{URL}?ket_qua={ls.ma_tai}").content.decode()
        self.assertNotIn(ls.ma_trich_xuat, trang)
        self.assertNotIn(duong_dan, trang)
        self.client.logout()
        self.assertEqual(self.client.get(duong_dan).status_code, 302)
        self.assertEqual(LichSuXuat.objects.get(pk=ls.pk).so_lan_tai, 1)  # chỉ lần tải của chủ sở hữu

    def test_mat_quyen_xuat_thi_khong_tai_lai_duoc(self):
        nguoi = tao_nguoi_dung("tamthoi", quyen=["xem_so_cbcs", "xuat_so_cbcs", "xem_moi_cong"])
        ls, _, _ = self.xuat_va_tai(nguoi=nguoi)
        nguoi.user_permissions.remove(*nguoi.user_permissions.filter(codename="xuat_so_cbcs"))
        nguoi = type(nguoi).objects.get(pk=nguoi.pk)
        self.client.force_login(nguoi)
        self.assertEqual(self.client.get(reverse("baocao:tai_tep", args=[ls.ma_tai])).status_code, 404)

    def test_ma_tai_khong_doan_duoc(self):
        a, _, _ = self.xuat_va_tai()
        b, _, _ = self.xuat_va_tai()
        self.assertEqual(a.ma_tai.version, 4)
        self.assertNotEqual(a.ma_tai, b.ma_tai)
        self.assertEqual(CanBo.objects.count(), 7)  # không có dữ liệu nào được tạo thêm khi xuất

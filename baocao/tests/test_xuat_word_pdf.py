"""Xuất Word (python-docx) và PDF (ReportLab): nội dung, bố cục, đối chiếu, quyền, khối ký, giới hạn."""
import datetime
import hashlib
import io
import re
import shutil
import subprocess
import tempfile
import time
import unicodedata
from pathlib import Path
from unittest import mock, skipUnless

from django.core.exceptions import ValidationError
from django.forms.models import model_to_dict
from django.test import TestCase
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Mm
from pypdf import PdfReader

from ravao.models import CanBo, DinhChinh, DonVi, Luot

from .. import bo_cuc
from ..models import CauHinhSo, LichSuXuat
from .du_lieu import DS01, phan_cong, t, tao_nguoi_dung

S1 = {"tu": "2026-10-25T00:00", "den": "2026-10-26T00:00"}
GIO_S1 = t("09:00 27/10")
SOFFICE = shutil.which("soffice") or shutil.which("libreoffice")

TIEU_DE_COT = {
    "cbcs": ["STT", "Mã CBCS", "Họ và tên", "Đơn vị / bộ phận", "Chiều lượt", "Ra", "Vào",
             "Thời lượng\n(giờ:phút)", "Lý do / nội dung", "Phương tiện đi kèm", "Người ghi nhận", "Trạng thái",
             "Ghi chú"],
    "khach": ["STT", "Khách", "Giấy tờ", "Mục đích", "Người / bộ phận tiếp", "Số thẻ", "Vào", "Ra",
              "Thời lượng\n(giờ:phút)", "Phương tiện", "Người ghi nhận", "Trạng thái", "Ghi chú"],
    "phuong-tien": ["STT", "Biển số", "Loại phương tiện", "Nguồn", "Người điều khiển", "Chiều lượt", "Ra", "Vào",
                    "Thời lượng\n(giờ:phút)", "Lý do / lệnh điều xe", "Người ghi nhận", "Trạng thái", "Ghi chú"],
}
THU_TU_S1 = {"cbcs": ["C3", "C8", "C1", "C6", "C2", "C5", "C4"], "khach": ["K1", "K3", "K4", "K2", "K5", "K6"],
             "phuong-tien": ["P1", "P3", "P5", "P2", "P4"]}


# ---------------------------------------------------------------------------
# Đọc tệp
# ---------------------------------------------------------------------------


def trang_pdf(noi_dung):
    reader = PdfReader(io.BytesIO(noi_dung))
    return reader, [unicodedata.normalize("NFC", p.extract_text() or "") for p in reader.pages]


def gon(chu):
    """Bỏ mọi khoảng trắng: chữ trích từ PDF bị ngắt dòng theo độ rộng ô."""
    return re.sub(r"\s+", "", chu)


def phong_pdf(reader):
    """{tên phông: có nhúng} cho mọi phông được trang tham chiếu."""
    ket_qua = {}
    for trang in reader.pages:
        phong = (trang.get("/Resources") or {}).get("/Font") or {}
        for ref in phong.values():
            f = ref.get_object()
            mo_ta = f.get("/FontDescriptor")
            mo_ta = mo_ta.get_object() if mo_ta else {}
            ket_qua[str(f.get("/BaseFont"))] = any(k in mo_ta for k in ("/FontFile", "/FontFile2", "/FontFile3"))
    return ket_qua


def bang_word(doc, o_dau="STT"):
    """Bảng đầu tiên có ô đầu tiên của hàng tiêu đề cuối là ``o_dau``."""
    for bang in doc.tables:
        for hang in bang.rows[:2]:
            if hang.cells[0].text == o_dau:
                return bang
    raise AssertionError("Không thấy bảng")


def chu_word(doc):
    phan = [p.text for p in doc.paragraphs]
    for bang in doc.tables:
        for hang in bang.rows:
            phan += [o.text for o in hang.cells]
    return "\n".join(phan)


def chi_so_hang_tieu_de(bang):
    return sum(1 for tr in bang._tbl.tr_lst if tr.find(qn("w:trPr")) is not None
               and tr.find(qn("w:trPr")).find(qn("w:tblHeader")) is not None)


class CoSo(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.ds = DS01()
        cls.cau_hinh = CauHinhSo.objects.create(
            co_quan_chu_quan="Cơ quan thử nghiệm", ten_don_vi="Đơn vị thử nghiệm", dia_danh="Hà Nội")

    def setUp(self):
        dong_ho = mock.patch("baocao.dong_ho.hien_tai", return_value=GIO_S1)
        dong_ho.start()
        self.addCleanup(dong_ho.stop)

    def xuat(self, so="cbcs", dinh_dang="pdf", nguoi=None, **tham_so):
        self.client.force_login(nguoi or self.ds.tb01)
        duong_dan = {"pdf": "xuat-pdf", "docx": "xuat-word"}[dinh_dang]
        return self.client.get(f"/so/{so}/{duong_dan}/", {**S1, **tham_so})

    def xuat_ok(self, so="cbcs", dinh_dang="pdf", nguoi=None, **tham_so):
        phan_hoi = self.xuat(so, dinh_dang, nguoi, **tham_so)
        self.assertEqual(phan_hoi.status_code, 200, phan_hoi.content[:800])
        if dinh_dang == "pdf":
            return phan_hoi, trang_pdf(phan_hoi.content)
        return phan_hoi, Document(io.BytesIO(phan_hoi.content))

    def pk(self, *ten):
        return [self.ds.l[x].pk for x in ten]


# ---------------------------------------------------------------------------
# Nội dung và bố cục
# ---------------------------------------------------------------------------


class NoiDungPdfTest(CoSo):
    def test_thong_tin_dau_so_va_sieu_du_lieu(self):
        phan_hoi, (reader, trang) = self.xuat_ok()
        ma = phan_hoi["X-Ma-Trich-Xuat"]
        self.assertEqual(phan_hoi["Content-Type"], "application/pdf")
        self.assertEqual(phan_hoi["Content-Disposition"], f'attachment; filename="SO01_CBCS_20261025-20261025_{ma}.pdf"')
        self.assertTrue(phan_hoi.content.startswith(b"%PDF-"))
        dau = trang[0]
        for chu in ("CƠ QUAN THỬ NGHIỆM", "ĐƠN VỊ THỬ NGHIỆM", "Mã sổ: SO-01", "SỔ THEO DÕI CBCS RA/VÀO",
                    "Từ 00:00 ngày 25/10/2026 đến hết ngày 25/10/2026",
                    "Số liệu tính đến: 00:00:00 ngày 26/10/2026",
                    f"Thời điểm trích xuất: 09:00:00 ngày 27/10/2026 (giờ máy chủ) · Mã trích xuất: {ma}",
                    "Người xuất: Nguyễn Văn X (tài khoản: trucban01) · Người lập sổ: Nguyễn Văn X"):
            self.assertIn(chu, dau)
        self.assertEqual(reader.metadata.subject, ma)
        self.assertEqual(reader.metadata.title, "SỔ THEO DÕI CBCS RA/VÀO — Từ 00:00 ngày 25/10/2026 đến hết ngày 25/10/2026")
        self.assertEqual(reader.metadata.author, "Nguyễn Văn X (tài khoản: trucban01)")

    def test_kho_giay_va_phong_nhung(self):
        for so, kich_thuoc in (("cbcs", (297, 210)), ("khach", (297, 210)), ("phuong-tien", (297, 210)),
                               ("tong-hop", (210, 297))):
            with self.subTest(so):
                _, (reader, _) = self.xuat_ok(so)
                for trang in reader.pages:
                    rong, cao = float(trang.mediabox.width) / 72 * 25.4, float(trang.mediabox.height) / 72 * 25.4
                    self.assertAlmostEqual(rong, kich_thuoc[0], delta=0.5)
                    self.assertAlmostEqual(cao, kich_thuoc[1], delta=0.5)
                phong = phong_pdf(reader)
                self.assertTrue(phong)
                self.assertTrue(all(phong.values()), f"Có phông không nhúng: {phong}")
                self.assertTrue(all("LiberationSerif" in ten for ten in phong), phong)

    def test_danh_so_trang_va_dau_trang_rut_gon(self):
        _, (_, trang) = self.xuat_ok()
        n = len(trang)
        self.assertGreater(n, 1)
        for i, chu in enumerate(trang, 1):
            self.assertIn(f"Trang {i}/{n}", chu)
            self.assertIn("Mẫu nghiệp vụ nội bộ", chu)
            if i > 1:
                self.assertIn("ĐƠN VỊ THỬ NGHIỆM — SỔ THEO DÕI CBCS RA/VÀO — 25/10/2026", chu)

    def test_luot_chua_hoan_tat_qua_nua_dem_va_tong(self):
        _, (_, trang) = self.xuat_ok()
        chu = "\n".join(trang)
        # C3 từ kỳ trước, C4 hoàn tất sau kỳ, C5 đang mở quá hạn.
        self.assertRegex(chu, r"23:30 24/10\s*Cổng 1\s*\(từ kỳ trước\)")
        self.assertRegex(chu, r"00:20 26/10\s*Cổng 1\s*\(sau kỳ\)")
        self.assertIn("≥ 08:00", chu)
        self.assertRegex(chu, r"Chưa về –\s*Quá hạn")
        self.assertIn("Hoàn tất sau", chu)
        self.assertRegex(chu, r"Chưa hoàn tất tại thời điểm chốt\s*2\s*lượt")
        self.assertIn(gon("Cân đối: Từ kỳ trước + Phát sinh = Hoàn tất + Chưa hoàn tất (1 + 6 = 5 + 2)"), gon(chu))
        self.assertIn("(từ kỳ trước) — sự kiện xảy ra trước kỳ.", chu)
        self.assertIn("≥ — thời lượng tính đến thời điểm chốt", chu)


class NoiDungWordTest(CoSo):
    def test_cau_truc_bang_tieu_de_lap_va_khong_ngat_dong(self):
        for so in ("cbcs", "khach", "phuong-tien"):
            with self.subTest(so):
                phan_hoi, doc = self.xuat_ok(so, "docx")
                bang = bang_word(doc)
                self.assertEqual([o.text for o in bang.rows[0].cells], TIEU_DE_COT[so])
                self.assertEqual(chi_so_hang_tieu_de(bang), 1)
                for tr in bang._tbl.tr_lst:
                    self.assertIsNotNone(tr.find(qn("w:trPr")).find(qn("w:cantSplit")))
                self.assertEqual(len(bang.rows) - 1, len(THU_TU_S1[so]))
                self.assertEqual(doc.core_properties.subject, phan_hoi["X-Ma-Trich-Xuat"])

    def test_section_dau_trang_chan_trang(self):
        _, doc = self.xuat_ok("cbcs", "docx")
        sec = doc.sections[0]
        self.assertAlmostEqual(sec.page_width, Mm(297), delta=Mm(0.1))
        self.assertAlmostEqual(sec.page_height, Mm(210), delta=Mm(0.1))
        self.assertEqual(sec.orientation, 1)  # WD_ORIENT.LANDSCAPE
        self.assertAlmostEqual(sec.left_margin, Mm(20), delta=Mm(0.1))
        self.assertTrue(sec.different_first_page_header_footer)
        self.assertEqual(sec.header.paragraphs[0].text, "ĐƠN VỊ THỬ NGHIỆM — SỔ THEO DÕI CBCS RA/VÀO — 25/10/2026")
        self.assertEqual(sec.first_page_header.paragraphs[0].text, "")
        for chan in (sec.footer, sec.first_page_footer):
            xml = chan._element.xml
            self.assertIn('w:instr="PAGE"', xml)
            self.assertIn('w:instr="NUMPAGES"', xml)
            self.assertIn("Bản xuất để biên tập", chan.paragraphs[0].text)

    def test_luot_chua_hoan_tat_in_dam_va_nhan_manh_qua_han(self):
        _, doc = self.xuat_ok("cbcs", "docx")
        bang = bang_word(doc)
        c5 = next(h for h in bang.rows[1:] if h.cells[0].text.endswith(f"#{self.ds.l['C5'].pk}"))
        gia_tri = [o.text for o in c5.cells]
        self.assertEqual(gia_tri[6], "—")
        self.assertEqual(gia_tri[7], "≥ 08:00")
        self.assertEqual(gia_tri[11], "Chưa về – Quá hạn")
        for o in c5.cells:
            for p in o.paragraphs:
                for r in p.runs:
                    if r.text.strip():  # bỏ qua run chỉ chứa ngắt dòng
                        self.assertTrue(r.bold, (o.text, r.text))
        run_qua_han = [r for r in c5.cells[11].paragraphs[0].runs if r.text == "Quá hạn"]
        self.assertTrue(run_qua_han and run_qua_han[0].underline)
        c1 = next(h for h in bang.rows[1:] if h.cells[0].text.endswith(f"#{self.ds.l['C1'].pk}"))
        self.assertFalse(any(r.bold for r in c1.cells[2].paragraphs[0].runs))

    def test_qua_nua_dem(self):
        _, doc = self.xuat_ok("cbcs", "docx")
        dong = {h.cells[0].text.split("\n")[1]: [o.text for o in h.cells] for h in bang_word(doc).rows[1:]}
        c3 = dong[f"#{self.ds.l['C3'].pk}"]
        c4 = dong[f"#{self.ds.l['C4'].pk}"]
        self.assertEqual((c3[5], c3[6], c3[7]), ("23:30 24/10\nCổng 1\n(từ kỳ trước)", "01:15 25/10\nCổng 1", "01:45"))
        self.assertEqual((c4[6], c4[7], c4[11]), ("00:20 26/10\nCổng 1\n(sau kỳ)", "01:35", "Hoàn tất sau kỳ"))
        tong = {h.cells[0].text: h.cells[1].text for h in bang_word(doc, "Chỉ tiêu").rows[1:]}
        self.assertEqual((tong["— Từ kỳ trước"], tong["— Phát sinh trong kỳ"], tong["Hoàn tất trong kỳ"],
                          tong["Chưa hoàn tất tại thời điểm chốt"], tong["— Hoàn tất sau kỳ"]), ("1", "6", "5", "2", "1"))

    def test_so_tong_hop(self):
        _, doc = self.xuat_ok("tong-hop", "docx")
        a = bang_word(doc, "")  # Phần A có tiêu đề hai tầng: ô đầu tầng trên để trống.
        self.assertEqual([o.text for o in a.rows[0].cells][1:9], ["SỐ LƯỢT"] * 7 + ["SỐ ĐỐI TƯỢNG"])
        self.assertEqual(chi_so_hang_tieu_de(a), 2)
        self.assertEqual([[o.text for o in h.cells][:10] for h in a.rows[2:]], [
            ["CBCS", "1", "6", "5", "2", "1", "—", "1", "6", "1"],
            ["Khách", "0", "6", "4", "2", "1", "1", "0", "4", "0"],
            ["Phương tiện", "0", "5", "4", "1", "0", "1", "0", "3", "0"],
        ])
        chu = chu_word(doc)
        self.assertNotIn("Tổng cộng", chu)
        self.assertIn("(Không có dòng tổng cộng giữa các nhóm.)", chu)
        self.assertEqual(doc.sections[0].orientation, 0)  # dọc


# ---------------------------------------------------------------------------
# Đối chiếu với bản ghi gốc, không thay đổi dữ liệu
# ---------------------------------------------------------------------------


class DoiChieuTest(CoSo):
    def test_ma_luot_doi_chieu_duoc_voi_ban_ghi_goc(self):
        for so in ("cbcs", "khach", "phuong-tien"):
            mong_doi = [f"#{pk}" for pk in self.pk(*THU_TU_S1[so])]
            with self.subTest(so, dinh_dang="docx"):
                phan_hoi, doc = self.xuat_ok(so, "docx")
                stt = [h.cells[0].text.split("\n") for h in bang_word(doc).rows[1:]]
                self.assertEqual([x[1] for x in stt], mong_doi)
                self.assertEqual([x[0] for x in stt], [str(i) for i in range(1, len(mong_doi) + 1)])
                self.assertEqual(LichSuXuat.objects.get(ma_trich_xuat=phan_hoi["X-Ma-Trich-Xuat"]).so_dong, len(mong_doi))
            with self.subTest(so, dinh_dang="pdf"):
                phan_hoi, (_, trang) = self.xuat_ok(so, "pdf")
                chu = "\n".join(trang)
                vi_tri = [chu.index(m) for m in mong_doi]
                self.assertEqual(vi_tri, sorted(vi_tri))  # đúng thứ tự
                ls = LichSuXuat.objects.get(ma_trich_xuat=phan_hoi["X-Ma-Trich-Xuat"])
                self.assertEqual((ls.dinh_dang, ls.so_dong, ls.ket_qua), ("pdf", len(mong_doi), "THANH_CONG"))
                self.assertEqual(ls.sha256, hashlib.sha256(phan_hoi.content).hexdigest())

    def test_ma_luot_nhieu_chu_so_khong_bi_ngat(self):
        l = self.ds.l["C1"]
        Luot.objects.create(
            id=98765432, nhom="CBCS", chieu="RV", can_bo=l.can_bo, mo_thoi_diem=t("12:00 25/10"), mo_cong=self.ds.c1,
            mo_ghi_nhan_luc=t("12:00 25/10"), mo_nguoi_ghi=self.ds.tb01)
        _, (_, trang) = self.xuat_ok("cbcs", "pdf")
        self.assertIn("#98765432", "\n".join(trang))  # nguyên vẹn trên một dòng
        _, doc = self.xuat_ok("cbcs", "docx")
        self.assertIn("#98765432", [h.cells[0].text.split("\n")[1] for h in bang_word(doc).rows[1:]])

    def test_phu_luc_dinh_chinh_co_ma_luot(self):
        _, doc = self.xuat_ok("cbcs", "docx")
        pl = bang_word(doc, "(n)")
        self.assertEqual([o.text for o in pl.rows[1].cells][:5],
                         ["(1)", "4", f"#{self.ds.l['C6'].pk}", "CB005 Hoàng Văn Đức", "Thời điểm sự kiện đóng"])

    def test_khong_thay_doi_du_lieu_goc(self):
        def anh_chup():
            return ([model_to_dict(l) for l in Luot.objects.order_by("pk")],
                    [model_to_dict(d) for d in DinhChinh.objects.order_by("pk")],
                    list(CanBo.objects.order_by("pk").values()))

        truoc = anh_chup()
        for so in ("cbcs", "khach", "phuong-tien", "tong-hop"):
            for dinh_dang in ("docx", "pdf"):
                self.xuat_ok(so, dinh_dang)
        self.xuat_ok("cbcs", "pdf", nguoi=self.ds.chihuy, hien_huy="on")
        self.assertEqual(anh_chup(), truoc)


# ---------------------------------------------------------------------------
# Nhiều trang
# ---------------------------------------------------------------------------


class NhieuTrangTest(CoSo):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        don_vi = DonVi.objects.first()
        can_bo = CanBo.objects.bulk_create([CanBo(ma=f"T{i:04d}", ho_ten=f"Cán bộ Thử Nghiệm {i}", don_vi=don_vi)
                                            for i in range(120)])
        bat_dau = t("00:00 25/10")
        Luot.objects.bulk_create([
            Luot(nhom="CBCS", chieu="RV", can_bo=can_bo[i % 120], mo_thoi_diem=bat_dau + datetime.timedelta(minutes=3 * i),
                 mo_cong=cls.ds.c1, mo_ghi_nhan_luc=bat_dau + datetime.timedelta(minutes=3 * i), mo_nguoi_ghi=cls.ds.tb01,
                 dong_thoi_diem=bat_dau + datetime.timedelta(minutes=3 * i + 45), dong_cong=cls.ds.c1,
                 dong_ghi_nhan_luc=bat_dau + datetime.timedelta(minutes=3 * i + 45), dong_nguoi_ghi=cls.ds.tb01,
                 ly_do="Công tác")
            for i in range(400)
        ])
        cls.so_luot = 400 + 7
        cls.ma_bang = {str(pk) for pk in Luot.objects.filter(can_bo__ma__startswith="T").values_list("pk", flat=True)}

    def test_pdf_tieu_de_bang_lap_moi_trang(self):
        _, (reader, trang) = self.xuat_ok()
        n = len(trang)
        self.assertGreaterEqual(n, 10)
        ma_luot = set(re.findall(r"#(\d+)", "\n".join(trang)))
        self.assertGreaterEqual(len(ma_luot), self.so_luot)
        co_bang = [i for i, chu in enumerate(trang) if self.ma_bang & set(re.findall(r"#(\d+)", chu))]
        self.assertGreater(len(co_bang), 8)
        for i in co_bang:
            self.assertIn("Mã CBCS", trang[i], f"trang {i + 1} thiếu hàng tiêu đề")
            self.assertIn("Họ và tên", trang[i])
            self.assertIn(f"Trang {i + 1}/{n}", trang[i])
        # Mỗi lượt xuất hiện đúng một lần: không dòng nào bị ngắt sang hai trang.
        dem = [ma for chu in trang for ma in re.findall(r"#(\d+)", chu) if ma in self.ma_bang]
        self.assertEqual(sorted(dem), sorted(self.ma_bang))

    def test_word_tieu_de_lap_va_so_dong(self):
        phan_hoi, doc = self.xuat_ok("cbcs", "docx")
        bang = bang_word(doc)
        self.assertEqual(len(bang.rows) - 1, self.so_luot)
        self.assertEqual(chi_so_hang_tieu_de(bang), 1)
        self.assertTrue(all(tr.trPr.find(qn("w:cantSplit")) is not None for tr in bang._tbl.tr_lst))

    @skipUnless(SOFFICE, "Không có LibreOffice để dàn trang thử tệp Word")
    def test_word_dan_trang_that_bang_libreoffice(self):
        """Dàn trang tệp Word bằng LibreOffice để xác nhận tiêu đề bảng lặp và số trang khi hiển thị thật."""
        phan_hoi, _ = self.xuat_ok("cbcs", "docx")
        with tempfile.TemporaryDirectory() as tam:
            tep = Path(tam) / "so.docx"
            tep.write_bytes(phan_hoi.content)
            subprocess.run([SOFFICE, "--headless", "--convert-to", "pdf", "--outdir", tam, str(tep)],
                           check=True, capture_output=True, timeout=180)
            _, trang = trang_pdf((Path(tam) / "so.pdf").read_bytes())
        n = len(trang)
        self.assertGreaterEqual(n, 10)
        co_bang = [i for i, chu in enumerate(trang) if self.ma_bang & set(re.findall(r"#(\d+)", chu))]
        self.assertGreater(len(co_bang), 8)
        for i in co_bang:
            self.assertIn("Mã CBCS", trang[i], f"trang {i + 1} thiếu hàng tiêu đề")
            self.assertIn(f"Trang {i + 1}/{n}", trang[i])
        for i in range(1, n):
            self.assertIn("ĐƠN VỊ THỬ NGHIỆM — SỔ THEO DÕI CBCS RA/VÀO", trang[i])


# ---------------------------------------------------------------------------
# Tiếng Việt, dữ liệu dài, không có dữ liệu
# ---------------------------------------------------------------------------


class TiengVietTest(CoSo):
    TEN = ["Nguyễn Thị Ánh Tuyết", "Ưng Hoàng Phúc", "Ông Ích Khiêm", "Bùi Thị Ơn", "Đặng Quốc Hưng",
           "Lương Thế Vinh", "Phạm Thị Ngọc Hà"]

    def test_ho_ten_va_nhan_co_dau(self):
        don_vi = DonVi.objects.create(ten="Phòng Hậu cần – Kỹ thuật")
        for i, ten in enumerate(self.TEN[:4] + self.TEN[5:]):
            cb = CanBo.objects.create(ma=f"VN{i}", ho_ten=unicodedata.normalize("NFC", ten), don_vi=don_vi)
            Luot.objects.create(nhom="CBCS", chieu="RV", can_bo=cb, mo_thoi_diem=t("12:00 25/10"), mo_cong=self.ds.c1,
                                mo_ghi_nhan_luc=t("12:00 25/10"), mo_nguoi_ghi=self.ds.tb01,
                                ly_do="Công tác", ly_do_chi_tiet="Đi Đà Lạt, Huế, Cần Thơ")
        _, doc = self.xuat_ok("cbcs", "docx", sap_xep="doi_tuong")
        ten_word = [h.cells[2].text for h in bang_word(doc).rows[1:]]
        _, (_, trang) = self.xuat_ok("cbcs", "pdf")
        chu_pdf = gon("\n".join(trang))
        for ten in self.TEN:
            self.assertIn(ten, ten_word)
            self.assertIn(gon(ten), chu_pdf)
        self.assertIn(gon("Công tác: Đi Đà Lạt, Huế, Cần Thơ"), chu_pdf)
        self.assertIn("Phòng Hậu cần – Kỹ thuật", chu_word(doc))
        # Sắp theo tên tiếng Việt.
        thu_tu = [x.split()[-1] for x in ten_word]
        self.assertLess(thu_tu.index("Hưng"), thu_tu.index("Khiêm"))
        self.assertLess(thu_tu.index("Khiêm"), thu_tu.index("Ơn"))
        self.assertLess(thu_tu.index("Ơn"), thu_tu.index("Phúc"))

    def test_ky_tu_dac_biet_khong_lam_hong_tep(self):
        Luot.objects.filter(pk=self.ds.l["C1"].pk).update(
            ghi_chu="<b>Không phải thẻ</b> & \"trích dẫn\" 'đơn' \x01 =1+1", ly_do_chi_tiet="A < B > C")
        _, doc = self.xuat_ok("cbcs", "docx")
        _, (_, trang) = self.xuat_ok("cbcs", "pdf")
        chu = gon("\n".join(trang))
        self.assertIn("<b>Không phải thẻ</b> & \"trích dẫn\" 'đơn' \ufffd =1+1", chu_word(doc))
        # PDF: phông nhúng không có U+FFFD nên ký tự đã thay được in là "□".
        self.assertIn(gon("<b>Không phải thẻ</b> & \"trích dẫn\" 'đơn' \u25a1 =1+1"), chu)
        self.assertIn(gon("A < B > C"), chu)


class DuLieuDaiTest(CoSo):
    def test_noi_dung_dai_in_day_du_o_phu_luc(self):
        dai = ("Nội dung ghi chú rất dài của trực ban về tình huống phát sinh tại cổng, có dấu tiếng Việt. " * 25).strip()
        Luot.objects.filter(pk=self.ds.l["C1"].pk).update(ghi_chu=dai)
        _, doc = self.xuat_ok("cbcs", "docx")
        o = next(h for h in bang_word(doc).rows[1:] if h.cells[0].text.endswith(f"#{self.ds.l['C1'].pk}")).cells[12]
        self.assertTrue(o.text.startswith(dai[:60]))
        self.assertTrue(o.text.endswith("… (đầy đủ: PL-1)"))
        doan = [p.text for p in doc.paragraphs]
        i = doan.index(f"PL-1 — dòng 3, mã lượt #{self.ds.l['C1'].pk}, cột Ghi chú:")
        self.assertEqual(doan[i + 1], dai)
        self.assertIn("PL-n — nội dung dài được in đầy đủ ở Phụ lục nội dung dài mục n.", doan)

        _, (_, trang) = self.xuat_ok("cbcs", "pdf")
        self.assertIn(gon(dai), gon("".join(trang)))

    def test_chuoi_rat_dai_khong_co_dau_cach(self):
        Luot.objects.filter(pk=self.ds.l["C1"].pk).update(ghi_chu="W" * 40000)
        _, doc = self.xuat_ok("cbcs", "docx")
        # Phụ lục in đủ 40.000 ký tự; ô trong bảng trích 100 ký tự đầu.
        self.assertEqual(sorted(len(x) for x in re.findall("W{50,}", chu_word(doc))), [100, 40000])
        _, (_, trang) = self.xuat_ok("cbcs", "pdf")
        # Chuỗi trải qua nhiều trang nên bị đầu/chân trang chia thành nhiều đoạn: so tổng độ dài.
        self.assertEqual(sum(len(x) for x in re.findall("W+", gon("\n".join(trang)))), 40000 + 100)
        self.assertGreater(len(trang), 5)

    def test_ten_va_co_quan_dai_trong_o(self):
        co_quan = "Công ty Cổ phần Đầu tư Xây dựng và Phát triển Hạ tầng Kỹ thuật Đô thị Miền Bắc, " * 2
        self.ds.l["K1"].khach.__class__.objects.filter(pk=self.ds.l["K1"].khach_id).update(co_quan=co_quan.strip())
        _, doc = self.xuat_ok("khach", "docx")
        o = bang_word(doc).rows[1].cells[1].text
        self.assertIn(co_quan.strip(), o)  # dưới ngưỡng phụ lục: in đầy đủ, tự xuống dòng
        self.xuat_ok("khach", "pdf")


class KhongCoDuLieuTest(CoSo):
    TRONG = {"tu": "2026-01-01T00:00", "den": "2026-01-02T00:00"}

    def test_so_chi_tiet_trong(self):
        for so in ("cbcs", "khach", "phuong-tien"):
            with self.subTest(so, dinh_dang="docx"):
                _, doc = self.xuat_ok(so, "docx", **self.TRONG)
                bang = bang_word(doc)
                self.assertEqual([o.text for o in bang.rows[0].cells], TIEU_DE_COT[so])
                self.assertEqual(len(bang.rows), 1)
                chu = chu_word(doc)
                self.assertIn("Không có lượt nào trong kỳ theo điều kiện đã chọn.", chu)
                self.assertIn("Không có đính chính.", chu)
                self.assertIn("— Hết sổ —", chu)
                tong = {h.cells[0].text: h.cells[1].text for h in bang_word(doc, "Chỉ tiêu").rows[1:]}
                self.assertEqual(tong["Số lượt trong sổ"], "0")
            with self.subTest(so, dinh_dang="pdf"):
                _, (_, trang) = self.xuat_ok(so, "pdf", **self.TRONG)
                self.assertIn("Không có lượt nào trong kỳ theo điều kiện đã chọn.", "\n".join(trang))
                self.assertNotIn("CHÚ GIẢI KÝ HIỆU", "\n".join(trang))

    def test_so_tong_hop_trong(self):
        _, doc = self.xuat_ok("tong-hop", "docx", **self.TRONG)
        chu = chu_word(doc)
        self.assertEqual(chu.count("Không có lượt chưa hoàn tất."), 3)
        _, (_, trang) = self.xuat_ok("tong-hop", "pdf", **self.TRONG)
        self.assertIn("D1. CBCS (0 lượt)", "\n".join(trang))


# ---------------------------------------------------------------------------
# Phân quyền, dữ liệu cá nhân
# ---------------------------------------------------------------------------


class PhanQuyenTest(CoSo):
    def test_khong_co_quyen(self):
        chi_xem = tao_nguoi_dung("xem01", quyen=["xem_so_cbcs", "xem_moi_cong"])
        khong_gi = tao_nguoi_dung("khonggi")
        for dinh_dang in ("docx", "pdf"):
            for nguoi, so in [(chi_xem, "cbcs"), (khong_gi, "cbcs"), (chi_xem, "khach"), (khong_gi, "tong-hop")]:
                with self.subTest(dinh_dang=dinh_dang, nguoi=nguoi.username, so=so):
                    self.assertEqual(self.xuat(so, dinh_dang, nguoi).status_code, 403)
            self.assertEqual(self.xuat("cbcs", dinh_dang, hien_huy="on").status_code, 403)
            self.assertEqual(self.xuat("khach", dinh_dang, du_lieu_day_du="on").status_code, 403)
        self.client.logout()
        self.assertEqual(self.client.get("/so/cbcs/xuat-pdf/", S1).status_code, 302)
        self.assertFalse(LichSuXuat.objects.exists())

    def test_che_so_giay_to_va_khong_co_so_dien_thoai(self):
        _, doc = self.xuat_ok("khach", "docx")
        _, (_, trang) = self.xuat_ok("khach", "pdf")
        for chu in (gon(chu_word(doc)), gon("\n".join(trang))):
            self.assertIn("CCCD********0111", chu)
            for bi_mat in ("001090000111", "0912345678", "001190000444"):
                self.assertNotIn(bi_mat, chu)
        _, doc = self.xuat_ok("khach", "docx", nguoi=self.ds.chihuy, du_lieu_day_du="on")
        self.assertIn("CCCD 001090000111", chu_word(doc))
        self.assertIn("Dữ liệu cá nhân: hiển thị đầy đủ", chu_word(doc))

    def test_gioi_han_theo_cong(self):
        tb = tao_nguoi_dung("cong2", quyen=["xem_so_cbcs", "xuat_so_cbcs"])
        phan_cong(tb, self.ds.c2)
        _, doc = self.xuat_ok("cbcs", "docx", nguoi=tb)
        self.assertEqual([h.cells[0].text.split("\n")[1] for h in bang_word(doc).rows[1:]],
                         [f"#{pk}" for pk in self.pk("C2", "C4")])

    def test_so_tong_hop_an_dinh_danh_neu_khong_co_quyen_nhom(self):
        lanh_dao = tao_nguoi_dung("lanhdao01", quyen=["xem_so_tong_hop", "xuat_so_tong_hop", "xem_moi_cong"])
        _, (_, trang) = self.xuat_ok("tong-hop", "pdf", nguoi=lanh_dao)
        chu = "\n".join(trang)
        for ten in ("Phạm Thị Dung", "Ông Quang", "Lý Thị Sen", "29X1-111.11"):
            self.assertNotIn(gon(ten), gon(chu))
        self.assertEqual(gon(chu).count("(khôngcóquyềnxemchitiết)"), 5)


# ---------------------------------------------------------------------------
# Khối ký theo cấu hình phê duyệt
# ---------------------------------------------------------------------------


class KhoiKyTest(CoSo):
    def test_chua_phe_duyet_thi_khong_in(self):
        for dinh_dang in ("docx", "pdf"):
            phan_hoi = self.xuat("cbcs", dinh_dang)
            noi_dung = (chu_word(Document(io.BytesIO(phan_hoi.content))) if dinh_dang == "docx"
                        else "\n".join(trang_pdf(phan_hoi.content)[1]))
            self.assertNotIn("NGƯỜI LẬP SỔ", noi_dung)
            self.assertNotIn("(Ký, ghi rõ họ tên)", noi_dung)
            self.assertIn("— Hết sổ —", noi_dung)

    def test_in_dung_cau_hinh_da_phe_duyet(self):
        CauHinhSo.objects.filter(pk=self.cau_hinh.pk).update(
            khoi_ky_da_phe_duyet=True, can_cu_phe_duyet="Biên bản họp giao ban ngày 20/10/2026",
            dien_san_ngay_ky=True,
            o_ky=[{"chuc_danh": "Người lập sổ", "nguoi_lap_so": True}, {"chuc_danh": "Trực ban trưởng"},
                  {"chuc_danh": "Thủ trưởng đơn vị", "ap_dung": ["SO-04"]}],
        )
        _, doc = self.xuat_ok("cbcs", "docx", nguoi_lap_so="trucban03")
        chu = chu_word(doc)
        for x in ("NGƯỜI LẬP SỔ", "TRỰC BAN TRƯỞNG", "Trần Văn Z", "Hà Nội, ngày 27 tháng 10 năm 2026"):
            self.assertIn(x, chu)
        self.assertNotIn("THỦ TRƯỞNG ĐƠN VỊ", chu)
        _, (_, trang) = self.xuat_ok("tong-hop", "pdf")
        chu = "\n".join(trang)
        self.assertIn("THỦ TRƯỞNG ĐƠN VỊ", chu)
        self.assertIn("NGƯỜI LẬP SỔ", chu)
        # Khối ký nằm cùng trang với thông tin trích xuất (BR-IN-09).
        trang_ky = [i for i, x in enumerate(trang) if "NGƯỜI LẬP SỔ" in x]
        self.assertEqual(len(trang_ky), 1)
        self.assertIn("THÔNG TIN TRÍCH XUẤT", trang[trang_ky[0]])

    def test_kiem_tra_cau_hinh(self):
        ch = CauHinhSo(ten_don_vi="X", khoi_ky_da_phe_duyet=True, can_cu_phe_duyet="")
        with self.assertRaises(ValidationError):
            ch.full_clean()
        ch.can_cu_phe_duyet = "Quyết định 1"
        ch.o_ky = [{"chuc_danh": ""}, {"chuc_danh": "A", "ap_dung": ["SO-09"]}, {"chuc_danh": "B", "la": 1}]
        with self.assertRaises(ValidationError) as loi:
            ch.full_clean()
        self.assertEqual(len(loi.exception.message_dict["o_ky"]), 3)


# ---------------------------------------------------------------------------
# Khổ giấy, giới hạn, lỗi
# ---------------------------------------------------------------------------


class KhoGiayTest(TestCase):
    def test_chon_theo_so_cot_va_do_rong(self):
        def bang(so_cot, rong):
            return bo_cuc.Bang(cot=[bo_cuc.CotIn(str(i), 1, rong) for i in range(so_cot)], dong=[])

        self.assertEqual(bo_cuc.chon_kho_giay(bang(5, 20)), ("A4", "doc"))
        self.assertEqual(bo_cuc.chon_kho_giay(bang(13, 18)), ("A4", "ngang"))
        self.assertEqual(bo_cuc.chon_kho_giay(bang(20, 18)), ("A3", "ngang"))
        self.assertEqual(bo_cuc.chon_kho_giay(bang(40, 18)), ("A3", "ngang"))  # rộng nhất có thể


class GioiHanVaLoiTest(CoSo):
    def test_ky_qua_dai_cho_ban_in(self):
        for dinh_dang in ("docx", "pdf"):
            phan_hoi = self.xuat("cbcs", dinh_dang, tu="2026-09-01T00:00", den="2026-10-26T00:00")
            self.assertEqual(phan_hoi.status_code, 400)
            self.assertIn("tối đa 31 ngày", phan_hoi.content.decode())
        # Sổ tổng hợp và Excel không bị giới hạn 31 ngày.
        self.assertEqual(self.xuat("tong-hop", "pdf", tu="2026-09-01T00:00", den="2026-10-26T00:00").status_code, 200)
        self.client.force_login(self.ds.tb01)
        self.assertEqual(self.client.get("/so/cbcs/xuat-excel/", {"tu": "2026-09-01T00:00",
                                                                  "den": "2026-10-26T00:00"}).status_code, 200)
        self.assertFalse(LichSuXuat.objects.filter(dinh_dang__in=["docx", "pdf"], ma_so="SO-01").exists())

    def test_vuot_so_dong_thi_tu_choi_khong_tao_tep_cut(self):
        with mock.patch("baocao.dich_vu.SO_DONG_TOI_DA_IN", 5):
            phan_hoi = self.xuat("cbcs", "pdf")
        self.assertEqual(phan_hoi.status_code, 422)
        self.assertIn("vượt giới hạn 5 dòng", phan_hoi.content.decode())
        self.assertNotEqual(phan_hoi["Content-Type"], "application/pdf")
        ls = LichSuXuat.objects.get()
        self.assertEqual((ls.ket_qua, ls.sha256, ls.dinh_dang), ("THAT_BAI", "", "pdf"))

    def test_loi_tao_tep_ghi_that_bai(self):
        with mock.patch("baocao.in_word.tao_docx", side_effect=RuntimeError("hỏng")):
            phan_hoi = self.xuat("cbcs", "docx")
        self.assertEqual(phan_hoi.status_code, 500)
        self.assertIn("Không tạo được tệp Word.", phan_hoi.content.decode())
        self.assertEqual(LichSuXuat.objects.get().ket_qua, "THAT_BAI")

    def test_thieu_phong_thi_khong_tao_pdf(self):
        with mock.patch("baocao.in_pdf.pdfmetrics.getRegisteredFontNames", return_value=[]), \
                mock.patch("baocao.in_pdf._thu_muc_phong", return_value=Path("/khong/ton/tai")):
            phan_hoi = self.xuat("cbcs", "pdf")
        self.assertEqual(phan_hoi.status_code, 500)
        self.assertIn("Thiếu phông chữ", phan_hoi.content.decode())
        self.assertEqual(LichSuXuat.objects.get().ket_qua, "THAT_BAI")


class HieuNangTest(CoSo):
    def test_5000_dong(self):
        don_vi = DonVi.objects.first()
        can_bo = CanBo.objects.bulk_create([CanBo(ma=f"H{i:04d}", ho_ten=f"Cán bộ {i}", don_vi=don_vi)
                                            for i in range(200)])
        bat_dau = t("00:00 01/10")
        Luot.objects.bulk_create([
            Luot(nhom="CBCS", chieu="RV", can_bo=can_bo[i % 200],
                 mo_thoi_diem=bat_dau + datetime.timedelta(minutes=8 * i), mo_cong=self.ds.c1,
                 mo_ghi_nhan_luc=bat_dau + datetime.timedelta(minutes=8 * i), mo_nguoi_ghi=self.ds.tb01,
                 dong_thoi_diem=bat_dau + datetime.timedelta(minutes=8 * i + 60), dong_cong=self.ds.c1,
                 dong_ghi_nhan_luc=bat_dau + datetime.timedelta(minutes=8 * i + 60), dong_nguoi_ghi=self.ds.tb01,
                 ly_do="Công tác")
            for i in range(4990)
        ])
        ky = {"tu": "2026-10-01T00:00", "den": "2026-11-01T00:00"}
        for dinh_dang in ("docx", "pdf"):
            bat_dau_do = time.monotonic()
            with mock.patch("baocao.dong_ho.hien_tai", return_value=t("09:00 02/11")):
                phan_hoi = self.xuat("cbcs", dinh_dang, **ky)
            thoi_gian = time.monotonic() - bat_dau_do
            self.assertEqual(phan_hoi.status_code, 200)
            ls = LichSuXuat.objects.get(ma_trich_xuat=phan_hoi["X-Ma-Trich-Xuat"])
            self.assertEqual(ls.so_dong, 4990 + 8)
            if dinh_dang == "pdf":
                so_trang = len(PdfReader(io.BytesIO(phan_hoi.content)).pages)
            else:
                so_trang = "?"
                self.assertEqual(len(bang_word(Document(io.BytesIO(phan_hoi.content))).rows) - 1, 4998)
            print(f"\n[hieu_nang] {dinh_dang}: {ls.so_dong} dòng, {len(phan_hoi.content) / 1024:.0f} KiB, "
                  f"{so_trang} trang, {thoi_gian:.1f} giây")
            self.assertLess(thoi_gian, 120)

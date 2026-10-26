"""Tạo tệp Excel cho bốn sổ (BR-XF-02, report_definitions.md §2–§6).

Nguyên tắc:
* Văn bản luôn được ghi dạng chuỗi; chuỗi có thể bị Excel hiểu là công thức
  được đánh dấu "quote prefix" và không bao giờ được ghi thành công thức.
* Mã, số giấy tờ, số thẻ, biển số là chuỗi định dạng ``@`` để giữ số 0 đầu.
* Không cắt dữ liệu: chuỗi vượt giới hạn ô của Excel làm việc xuất thất bại.
* Giá trị tĩnh, không dùng công thức.
"""
import datetime
import io
from copy import copy
from dataclasses import dataclass

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties
from django.utils import timezone

from ravao.models import Luot, chuan_hoa_ma

from . import quyen, so_lieu

GIOI_HAN_KY_TU_O = 32767
GIOI_HAN_DONG = 1_048_576
KY_TU_CONG_THUC = ("=", "+", "-", "@", "\t", "\r", "\n", "＝", "＋", "－", "＠")

DINH_DANG_NGAY_GIO = "dd/mm/yyyy hh:mm:ss"
DINH_DANG_NGAY = "dd/mm/yyyy"
DINH_DANG_THOI_LUONG = "[h]:mm"
DINH_DANG_THOI_LUONG_TOI_THIEU = '"≥ "[h]:mm'
DINH_DANG_VAN_BAN = "@"

MM_SANG_INCH = 1 / 25.4

CHU_GIAI = [
    "ĐC(n) — dòng có đính chính, xem sheet Dinh_chinh mục (n).",
    "GBS — sự kiện được ghi bổ sung muộn hơn ngưỡng quy định.",
    "Chồng lượt — đối tượng có lượt mở chồng thời gian.",
    "Lỗi thời gian — thời điểm đóng sớm hơn thời điểm mở; không tính thời lượng.",
    "(từ kỳ trước) — sự kiện xảy ra trước kỳ; (sau kỳ) — sự kiện xảy ra sau thời điểm chốt.",
    "Thời lượng có “Là thời lượng tối thiểu = Có” là thời lượng tính đến thời điểm chốt của lượt chưa hoàn tất.",
]

_VIEN = Side(style="thin", color="000000")
_KHUNG = Border(left=_VIEN, right=_VIEN, top=_VIEN, bottom=_VIEN)
_NEN_TIEU_DE = PatternFill("solid", fgColor="D9D9D9")


class LoiDuLieuExcel(Exception):
    """Dữ liệu không ghi được vào Excel mà không làm mất nội dung."""


# ---------------------------------------------------------------------------
# Ghi ô an toàn
# ---------------------------------------------------------------------------


def van_ban_an_toan(gia_tri):
    """Chuẩn hoá chuỗi để ghi vào ô. Trả về (chuỗi, có_nguy_cơ_công_thức)."""
    chuoi = str(gia_tri)
    # Ký tự điều khiển không hợp lệ trong XML được thay bằng U+FFFD để
    # người đọc thấy có ký tự đã bị thay, thay vì âm thầm xoá.
    chuoi = ILLEGAL_CHARACTERS_RE.sub("�", chuoi)
    if len(chuoi) > GIOI_HAN_KY_TU_O:
        raise LoiDuLieuExcel(
            f"Nội dung dài {len(chuoi)} ký tự, vượt giới hạn {GIOI_HAN_KY_TU_O} ký tự của một ô Excel."
        )
    return chuoi, chuoi.startswith(KY_TU_CONG_THUC)


def ghi_o(ws, hang, cot, gia_tri, kieu="text", dam=False, khung=False, xuong_dong=False):
    o = ws.cell(row=hang, column=cot)
    nguy_co = False
    if gia_tri is None or gia_tri == "":
        o.value = None
        if kieu in ("text", "ma"):
            o.number_format = DINH_DANG_VAN_BAN
    elif kieu in ("text", "ma"):
        chuoi, nguy_co = van_ban_an_toan(gia_tri)
        o.value = chuoi
        o.data_type = "s"  # openpyxl tự coi chuỗi bắt đầu bằng "=" là công thức.
        o.number_format = DINH_DANG_VAN_BAN
        o.quotePrefix = nguy_co
    elif kieu == "ngay_gio":
        o.value = timezone.localtime(gia_tri).replace(tzinfo=None)
        o.number_format = DINH_DANG_NGAY_GIO
    elif kieu == "ngay":
        o.value = gia_tri
        o.number_format = DINH_DANG_NGAY
    elif kieu in ("thoi_luong", "thoi_luong_toi_thieu"):
        o.value = gia_tri
        # Thời lượng tính đến chốt của lượt chưa hoàn tất hiển thị "≥ h:mm" nhưng vẫn là giá trị số.
        o.number_format = DINH_DANG_THOI_LUONG_TOI_THIEU if kieu == "thoi_luong_toi_thieu" else DINH_DANG_THOI_LUONG
    elif kieu == "so":
        o.value = gia_tri
        o.number_format = "#,##0"
    else:
        raise ValueError(kieu)
    # Định dạng giống nhau được dựng một lần cho mỗi workbook rồi sao chép, vì
    # dựng và băm đối tượng định dạng cho từng ô chiếm phần lớn thời gian xuất.
    bo_dem = ws.parent.__dict__.setdefault("_dinh_dang_da_dung", {})
    khoa = (kieu if o.value is not None else "rong-" + kieu, dam, khung, xuong_dong, nguy_co)
    if khoa in bo_dem:
        o._style = copy(bo_dem[khoa])
        return o
    if dam:
        o.font = Font(bold=True)
    if khung:
        o.border = _KHUNG
    o.alignment = Alignment(vertical="top", wrap_text=xuong_dong)
    bo_dem[khoa] = copy(o._style)
    return o


def do_dai_hien_thi(gia_tri, kieu):
    if gia_tri is None or gia_tri == "":
        return 0
    if kieu == "ngay_gio":
        return 19
    if kieu == "ngay":
        return 10
    if kieu in ("thoi_luong", "thoi_luong_toi_thieu"):
        return 9
    return max(len(dong) for dong in str(gia_tri).split("\n"))


# ---------------------------------------------------------------------------
# Thông tin chung của một lần xuất
# ---------------------------------------------------------------------------


@dataclass
class ThongTinXuat:
    ma_trich_xuat: str
    thoi_diem_trich_xuat: datetime.datetime
    nguoi_xuat: str  # "Họ tên (tài khoản: x)"
    nguoi_lap_so: str  # rỗng = để trống, ký tay
    cau_hinh: object
    du_lieu_ca_nhan_day_du: bool = False


def _dt(gia_tri):
    return timezone.localtime(gia_tri).strftime("%H:%M:%S ngày %d/%m/%Y")


def mo_ta_chot(ky, chot):
    chu = f"Số liệu tính đến: {_dt(chot)}"
    return chu + (" (kỳ chưa kết thúc)" if chot < ky.den else "")


def mo_ta_cong(bo_loc):
    return "Điểm kiểm soát: " + (", ".join(c.ten for c in bo_loc.cong) if bo_loc.cong else "Tất cả các cổng")


def _ghi_dau_so(ws, dinh_nghia, ky, chot, bo_loc, tt, so_cot, dong_them=()):
    """Khối đầu sổ (report_definitions.md §2.3). Trả về hàng kế tiếp còn trống."""
    ch = tt.cau_hinh
    cot_phai = max(so_cot - 2, 2)
    ghi_o(ws, 1, 1, ch.co_quan_chu_quan.upper() if ch.co_quan_chu_quan else "")
    ghi_o(ws, 1, cot_phai, f"Mã sổ: {dinh_nghia.ma}")
    ghi_o(ws, 2, 1, ch.ten_don_vi.upper(), dam=True)
    ghi_o(ws, 2, cot_phai, "Mẫu nội bộ — phiên bản 1.0")
    o = ghi_o(ws, 3, 1, dinh_nghia.ten, dam=True)
    o.font = Font(bold=True, size=14)
    ghi_o(ws, 4, 1, ky.mo_ta())
    loc = bo_loc.mo_ta()
    ghi_o(ws, 5, 1, mo_ta_cong(bo_loc) + "    Điều kiện lọc: " + ("; ".join(loc) if loc else "Không"))
    ghi_o(ws, 6, 1, mo_ta_chot(ky, chot))
    ghi_o(ws, 7, 1, f"Thời điểm trích xuất: {_dt(tt.thoi_diem_trich_xuat)} (giờ máy chủ)    "
                    f"Mã trích xuất: {tt.ma_trich_xuat}")
    ghi_o(ws, 8, 1, f"Người xuất: {tt.nguoi_xuat}    Người lập sổ: {tt.nguoi_lap_so or '(để trống — ký tay)'}")
    hang = 9
    them = list(dong_them)
    if bo_loc.co_loc:
        them.insert(0, "Sổ đã lọc — không phải toàn bộ dữ liệu của kỳ.")
    if bo_loc.hien_huy:
        them.append("Có hiển thị bản ghi đã huỷ (không tính vào số liệu).")
    for chu in them:
        o = ghi_o(ws, hang, 1, chu)
        o.font = Font(italic=True)
        hang += 1
    return hang + 1


def _thiet_lap_in(ws, dinh_nghia, tt, ngang, hang_tieu_de=None, so_cot=None, hang_cuoi=None):
    """Khổ A4, hướng, lề, vừa chiều ngang, tiêu đề lặp, chân trang (BR-IN-02, BR-IN-05)."""
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE if ngang else ws.ORIENTATION_PORTRAIT
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins.left = 20 * MM_SANG_INCH
    ws.page_margins.right = ws.page_margins.top = ws.page_margins.bottom = 15 * MM_SANG_INCH
    ws.page_margins.header = ws.page_margins.footer = 6 * MM_SANG_INCH
    ws.print_options.horizontalCentered = True
    if hang_tieu_de:
        ws.print_title_rows = f"{hang_tieu_de}:{hang_tieu_de}"
    if so_cot and hang_cuoi:
        ws.print_area = f"A1:{get_column_letter(so_cot)}{hang_cuoi}"
    chan = (f"{dinh_nghia.ma} · Mã trích xuất {tt.ma_trich_xuat} · Trích xuất lúc "
            f"{timezone.localtime(tt.thoi_diem_trich_xuat):%H:%M:%S %d/%m/%Y} · Mẫu nghiệp vụ nội bộ · "
            "Trang &P/&N")
    ws.oddFooter.center.text = chan
    ws.oddFooter.center.size = 8
    ws.evenFooter.center.text = chan


@dataclass
class Cot:
    tieu_de: str
    lay: object  # hàm(dong) -> giá trị
    kieu: str = "text"
    rong_toi_thieu: int = 6
    rong_toi_da: int = 40
    in_an: bool = True  # False: cột chỉ có trong Excel, nằm ngoài vùng in (report_definitions.md §3.4)
    kieu_dong: object = None  # hàm(dong) -> kiểu, khi kiểu thay đổi theo dòng


def _ghi_bang(ws, hang_dau, cac_cot, dong, bo_loc_excel=True):
    """Ghi bảng có tiêu đề, khung, bộ lọc, cố định tiêu đề. Trả về hàng cuối."""
    if hang_dau + len(dong) > GIOI_HAN_DONG:
        raise LoiDuLieuExcel("Số dòng vượt giới hạn của một sheet Excel. Hãy thu hẹp kỳ hoặc bộ lọc.")
    rong = [do_dai_hien_thi(c.tieu_de, "text") for c in cac_cot]
    for j, c in enumerate(cac_cot, 1):
        o = ghi_o(ws, hang_dau, j, c.tieu_de, dam=True, khung=True, xuong_dong=True)
        o.fill = _NEN_TIEU_DE
        o.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    hang = hang_dau
    for d in dong:
        hang += 1
        dam = getattr(d, "trang_thai", None) == so_lieu.MO  # Lượt chưa hoàn tất in đậm (BR-TT-04).
        for j, c in enumerate(cac_cot, 1):
            gia_tri = c.lay(d)
            kieu = c.kieu_dong(d) if c.kieu_dong else c.kieu
            kieu = "text" if isinstance(gia_tri, str) and kieu == "so" else kieu
            ghi_o(ws, hang, j, gia_tri, kieu, dam=dam, khung=True, xuong_dong=kieu == "text")
            rong[j - 1] = max(rong[j - 1], do_dai_hien_thi(gia_tri, kieu))
    for j, c in enumerate(cac_cot, 1):
        ws.column_dimensions[get_column_letter(j)].width = min(max(rong[j - 1] + 2, c.rong_toi_thieu), c.rong_toi_da)
    if bo_loc_excel and dong:
        ws.auto_filter.ref = f"A{hang_dau}:{get_column_letter(len(cac_cot))}{hang}"
    ws.freeze_panes = ws.cell(row=hang_dau + 1, column=1)
    return hang


def _ghi_cap(ws, hang, nhan, gia_tri, kieu="so", cot_nhan=2, cot_gia_tri=None, don_vi=""):
    ghi_o(ws, hang, cot_nhan, nhan)
    cot_gia_tri = cot_gia_tri or cot_nhan + 4
    ghi_o(ws, hang, cot_gia_tri, gia_tri, kieu)
    if don_vi:
        ghi_o(ws, hang, cot_gia_tri + 1, don_vi)


def _ghi_khoi_ky(ws, hang, tt, cot_phai):
    ch = tt.cau_hinh
    dia_danh = ch.dia_danh or "..........."
    ghi_o(ws, hang, cot_phai, f"{dia_danh}, ngày ..... tháng ..... năm ..........")
    hang += 1
    ghi_o(ws, hang, 2, "NGƯỜI LẬP SỔ", dam=True)
    ghi_o(ws, hang, cot_phai, (ch.chuc_danh_xac_nhan or "").upper(), dam=True)
    hang += 1
    ghi_o(ws, hang, 2, "(Ký, ghi rõ họ tên)")
    ghi_o(ws, hang, cot_phai, "(Ký, ghi rõ họ tên)")
    hang += 4
    ghi_o(ws, hang, 2, tt.nguoi_lap_so or "...............................")
    ghi_o(ws, hang, cot_phai, "...............................")
    hang += 2
    ghi_o(ws, hang, 2, "— Hết sổ —")
    return hang


# ---------------------------------------------------------------------------
# Cột của từng sổ (report_definitions.md §3.4, §4.4, §5.4 — bản Excel tách ô)
# ---------------------------------------------------------------------------


def che_so_giay_to(so):
    """Giữ 4 ký tự cuối (BR-BM-01)."""
    if not so:
        return so
    return "*" * max(len(so) - 4, 0) + so[-4:] if len(so) > 4 else "*" * len(so)


def _sap_cot(cot):
    """Cột in đứng trước, cột chỉ có trong Excel đứng sau để vùng in liền một khối."""
    return [c for c in cot if c.in_an] + [c for c in cot if not c.in_an]


def so_cot_in(cot):
    return sum(c.in_an for c in cot)


def _cot_su_kien(loai, nhan):
    return [
        Cot(f"{nhan} – thời điểm", lambda d: d.su_kien(loai).thoi_diem, "ngay_gio", 19),
        Cot(f"{nhan} – cổng", lambda d: d.su_kien(loai).cong.ten if d.su_kien(loai).cong else ""),
        Cot(f"{nhan} – chú thích kỳ", lambda d: d.su_kien(loai).chu_thich, rong_toi_thieu=10),
    ]


def _cot_nguoi_ghi(loai, nhan):
    return [
        Cot(f"Người ghi {nhan}", lambda d: so_lieu.nguoi_ghi_hien_thi(d.su_kien(loai).nguoi_ghi)),
        Cot(f"Ghi nhận {nhan} lúc", lambda d: d.su_kien(loai).ghi_nhan_luc, "ngay_gio", 19, in_an=False),
    ]


def _cot_cuoi():
    return [
        Cot("Thời lượng (giờ:phút)", lambda d: d.thoi_luong, "thoi_luong", 10,
            kieu_dong=lambda d: "thoi_luong_toi_thieu" if d.thoi_luong_toi_thieu else "thoi_luong"),
        Cot("Là thời lượng tối thiểu", lambda d: "Có" if d.thoi_luong_toi_thieu else "", in_an=False),
    ]


def _cot_trang_thai():
    return [
        Cot("Mã trạng thái", lambda d: d.trang_thai, in_an=False),
        Cot("Trạng thái", lambda d: d.nhan_trang_thai, rong_toi_thieu=12),
        Cot("Ghi chú", lambda d: d.ky_hieu, rong_toi_thieu=12, rong_toi_da=60),
        Cot("Mã lượt", lambda d: d.luot.pk, "so", in_an=False),
    ]


def _stt(d):
    return d.stt if d.stt is not None else "—"


def _ten_nguon(pt):
    if pt.nguon == pt.Nguon.CBCS and pt.nguon_can_bo:
        return f"CBCS — {pt.nguon_can_bo.ho_ten}"
    if pt.nguon == pt.Nguon.KHACH and pt.nguon_khach:
        return f"Khách — {pt.nguon_khach.ho_ten}"
    if pt.nguon == pt.Nguon.KHAC and pt.nguon_ghi_chu:
        return f"Khác — {pt.nguon_ghi_chu}"
    return pt.get_nguon_display()


def cac_cot(nhom, du_lieu_ca_nhan_day_du):
    stt = Cot("STT", _stt, "so", 5, 6)
    if nhom == Luot.Nhom.CBCS:
        return _sap_cot([
            stt,
            Cot("Mã CBCS", lambda d: d.luot.can_bo.ma, "ma"),
            Cot("Họ và tên", lambda d: d.luot.can_bo.ho_ten, rong_toi_thieu=18),
            Cot("Cấp bậc", lambda d: d.luot.can_bo.cap_bac, in_an=False),
            Cot("Chức vụ", lambda d: d.luot.can_bo.chuc_vu, in_an=False),
            Cot("Đơn vị / bộ phận", lambda d: d.luot.can_bo.don_vi.ten, rong_toi_thieu=14),
            Cot("Chiều lượt", lambda d: d.luot.get_chieu_display()),
            *_cot_su_kien("RA", "Ra"),
            *_cot_su_kien("VAO", "Vào"),
            *_cot_cuoi(),
            Cot("Lý do", lambda d: d.luot.ly_do),
            Cot("Nội dung", lambda d: d.luot.ly_do_chi_tiet, rong_toi_da=50, in_an=False),
            Cot("Hạn về", lambda d: d.luot.han_du_kien, "ngay_gio", 19, in_an=False),
            Cot("Phương tiện đi kèm", lambda d: d.luot.phuong_tien_di_kem.hien_thi_bien_so
                if d.luot.phuong_tien_di_kem else "", "ma"),
            *_cot_nguoi_ghi("RA", "Ra"),
            *_cot_nguoi_ghi("VAO", "Vào"),
            *_cot_trang_thai(),
        ])
    if nhom == Luot.Nhom.KHACH:
        def so_giay_to(d):
            khach = d.luot.khach
            if not khach.so_giay_to:
                return "Không xuất trình"
            # Che trên số đã chuẩn hoá để khoảng trắng, dấu chấm không làm lộ thêm chữ số.
            return khach.so_giay_to if du_lieu_ca_nhan_day_du else che_so_giay_to(khach.so_giay_to_chuan)

        return _sap_cot([
            stt,
            Cot("Họ và tên", lambda d: d.luot.khach.ho_ten, rong_toi_thieu=18),
            Cot("Cơ quan / địa chỉ", lambda d: d.luot.khach.co_quan, rong_toi_thieu=14, rong_toi_da=50),
            Cot("Đoàn", lambda d: d.luot.ma_doan, "ma"),
            Cot("Loại giấy tờ", lambda d: d.luot.khach.loai_giay_to, in_an=False),
            Cot("Số giấy tờ", so_giay_to, "ma", 14),
            Cot("Mục đích", lambda d: d.luot.muc_dich, rong_toi_thieu=14, rong_toi_da=50),
            Cot("Người / bộ phận tiếp", lambda d: d.luot.nguoi_tiep, rong_toi_thieu=14),
            Cot("Số thẻ", lambda d: d.luot.so_the, "ma"),
            *_cot_su_kien("VAO", "Vào"),
            *_cot_su_kien("RA", "Ra"),
            *_cot_cuoi(),
            Cot("Giờ ra dự kiến", lambda d: d.luot.han_du_kien, "ngay_gio", 19, in_an=False),
            Cot("Phương tiện", lambda d: d.luot.phuong_tien_di_kem.hien_thi_bien_so
                if d.luot.phuong_tien_di_kem else "", "ma"),
            *_cot_nguoi_ghi("VAO", "Vào"),
            *_cot_nguoi_ghi("RA", "Ra"),
            *_cot_trang_thai(),
        ])
    return _sap_cot([
        stt,
        Cot("Biển số", lambda d: d.luot.phuong_tien.hien_thi_bien_so, "ma", 12),
        Cot("Biển số chuẩn hoá", lambda d: d.luot.phuong_tien.bien_so_chuan, "ma", in_an=False),
        Cot("Loại phương tiện", lambda d: d.luot.phuong_tien.loai),
        Cot("Nguồn", lambda d: _ten_nguon(d.luot.phuong_tien), rong_toi_thieu=12),
        Cot("Người điều khiển", lambda d: so_lieu.ten_can_bo(d.luot.nguoi_dieu_khien_can_bo)
            or d.luot.nguoi_dieu_khien, rong_toi_thieu=16),
        Cot("Chiều lượt", lambda d: d.luot.get_chieu_display()),
        *_cot_su_kien("RA", "Ra"),
        *_cot_su_kien("VAO", "Vào"),
        *_cot_cuoi(),
        Cot("Lý do", lambda d: d.luot.ly_do, rong_toi_da=50),
        Cot("Lệnh điều xe", lambda d: d.luot.lenh_dieu_xe, "ma", in_an=False),
        Cot("Hạn về", lambda d: d.luot.han_du_kien, "ngay_gio", 19, in_an=False),
        *_cot_nguoi_ghi("RA", "Ra"),
        *_cot_nguoi_ghi("VAO", "Vào"),
        *_cot_trang_thai(),
    ])


# ---------------------------------------------------------------------------
# Sổ chi tiết
# ---------------------------------------------------------------------------


_NHAN_DOI_TUONG = {
    Luot.Nhom.CBCS: ("Số CBCS duy nhất", "đối tượng"),
    Luot.Nhom.KHACH: ("Số khách có định danh duy nhất", "đối tượng"),
    Luot.Nhom.PHUONG_TIEN: ("Số phương tiện có biển duy nhất", "đối tượng"),
}
_NHAN_KHONG_DINH_DANH = {
    Luot.Nhom.KHACH: "Số lượt khách không định danh (không cộng vào số khách duy nhất)",
    Luot.Nhom.PHUONG_TIEN: "Số lượt phương tiện không biển (không cộng vào số phương tiện duy nhất)",
}


def _dong_chi_tieu(ct, nhom):
    """Các dòng của khối "Tổng hợp số liệu" (report_definitions.md §2.7)."""
    dong = [
        ("Số lượt trong sổ", ct["so_luot"], "lượt"),
        ("— Từ kỳ trước", ct["tu_ky_truoc"], "lượt"),
        ("— Phát sinh trong kỳ", ct["phat_sinh"], "lượt"),
        ("Hoàn tất trong kỳ", ct["hoan_tat"], "lượt"),
        ("Chưa hoàn tất tại thời điểm chốt", ct["chua_hoan_tat"], "lượt"),
        ("— Đang mở", ct["dang_mo"], "lượt"),
        ("— Hoàn tất sau kỳ", ct["hoan_tat_sau_ky"], "lượt"),
        ("— Trong đó quá hạn", ct["qua_han"], "lượt"),
        (_NHAN_DOI_TUONG[nhom][0], ct["doi_tuong_duy_nhat"], "đối tượng"),
    ]
    if nhom in _NHAN_KHONG_DINH_DANH:
        dong.append((_NHAN_KHONG_DINH_DANH[nhom], ct["khong_dinh_danh"], "lượt"))
    if nhom == Luot.Nhom.KHACH:
        dong += [("Số đoàn", ct["so_doan"], "đoàn"),
                 ("Thẻ khách chưa thu hồi tại thời điểm chốt", ct["the_chua_thu_hoi"], "thẻ")]
    for khoa, tieu_de in (("theo_chieu", "Theo chiều lượt"), ("theo_ly_do", "Theo lý do"),
                          ("theo_loai", "Theo loại phương tiện"), ("theo_nguon", "Theo nguồn")):
        if khoa in ct:
            dong.append((tieu_de, None, ""))
            dong += [(f"— {k}", v, "lượt") for k, v in ct[khoa].items()]
    dong += [
        ("Số lượt có đính chính", ct["luot_dinh_chinh"], "lượt"),
        ("Số bản ghi đã huỷ trong kỳ (không tính vào các chỉ tiêu trên)", ct["ban_ghi_huy"], "bản ghi"),
    ]
    return dong


def _doi_tuong_dinh_chinh(d):
    l = d.luot
    if l.nhom == Luot.Nhom.CBCS:
        return f"{l.can_bo.ma} {l.can_bo.ho_ten}"
    if l.nhom == Luot.Nhom.KHACH:
        return l.khach.ho_ten
    return l.phuong_tien.hien_thi_bien_so


def _nhan_truong(truong):
    if truong == "khach.so_giay_to":
        return "Số giấy tờ"
    if truong == "phuong_tien.bien_so":
        return "Biển số"
    try:
        return str(Luot._meta.get_field(truong).verbose_name).capitalize()
    except Exception:
        return truong


def _ghi_sheet_dinh_chinh(wb, ket_qua, tt):
    ws = wb.create_sheet("Dinh_chinh")
    ghi_o(ws, 1, 1, f"PHỤ LỤC ĐÍNH CHÍNH — {ket_qua.dinh_nghia.ma}", dam=True)
    ghi_o(ws, 2, 1, ket_qua.ky.mo_ta())

    def gia_tri(dc, cu):
        v = dc.gia_tri_cu if cu else dc.gia_tri_moi
        if dc.la_du_lieu_ca_nhan and not tt.du_lieu_ca_nhan_day_du:
            return che_so_giay_to(chuan_hoa_ma(v))
        return v

    cot = [
        Cot("(n)", lambda m: m[0], "so"),
        Cot("STT dòng", lambda m: m[1] if m[1] is not None else "— (đã huỷ)", "so"),
        Cot("Đối tượng", lambda m: _doi_tuong_dinh_chinh(m[2]), rong_toi_thieu=16),
        Cot("Trường", lambda m: _nhan_truong(m[3].truong)),
        Cot("Giá trị cũ", lambda m: gia_tri(m[3], True), "text", 12, 50),
        Cot("Giá trị mới", lambda m: gia_tri(m[3], False), "text", 12, 50),
        Cot("Lý do", lambda m: m[3].ly_do, rong_toi_da=50),
        Cot("Người đính chính", lambda m: _ho_ten_tai_khoan(m[3].nguoi)),
        Cot("Thời điểm đính chính", lambda m: m[3].thoi_diem, "ngay_gio", 19),
    ]
    if ket_qua.phu_luc:
        hang = _ghi_bang(ws, 4, cot, ket_qua.phu_luc)
    else:
        ghi_o(ws, 4, 1, "Không có đính chính.")
        hang = 4
    _thiet_lap_in(ws, ket_qua.dinh_nghia, tt, ngang=True, hang_tieu_de=4 if ket_qua.phu_luc else None,
                  so_cot=len(cot), hang_cuoi=hang)


def _ho_ten_tai_khoan(nguoi):
    return f"{quyen.ho_ten(nguoi)} ({nguoi.get_username()})"


def _ghi_sheet_thong_tin(wb, dinh_nghia, ky, chot, bo_loc, tt, so_dong, ghi_chu_them=()):
    ws = wb.create_sheet("Thong_tin_trich_xuat")
    loc = bo_loc.mo_ta()
    muc = [
        ("Mã sổ", dinh_nghia.ma, "text"),
        ("Tên sổ", dinh_nghia.ten, "text"),
        ("Mã trích xuất", tt.ma_trich_xuat, "ma"),
        ("Định dạng", "Excel (.xlsx)", "text"),
        ("Kỳ", ky.mo_ta(), "text"),
        ("Kỳ — từ (gồm)", ky.tu, "ngay_gio"),
        ("Kỳ — đến (không gồm)", ky.den, "ngay_gio"),
        ("Số liệu tính đến (thời điểm chốt)", chot, "ngay_gio"),
        ("Thời điểm trích xuất (giờ máy chủ)", tt.thoi_diem_trich_xuat, "ngay_gio"),
        ("Người xuất", tt.nguoi_xuat, "text"),
        ("Người lập sổ", tt.nguoi_lap_so or "(để trống — ký tay)", "text"),
        ("Điểm kiểm soát", mo_ta_cong(bo_loc).split(": ", 1)[1], "text"),
        ("Điều kiện lọc", "; ".join(loc) if loc else "Không", "text"),
        ("Hiện bản ghi đã huỷ", "Có" if bo_loc.hien_huy else "Không", "text"),
        ("Dữ liệu cá nhân", "Hiển thị đầy đủ" if tt.du_lieu_ca_nhan_day_du else "Số giấy tờ đã che", "text"),
        ("Số dòng dữ liệu", so_dong, "so"),
    ]
    for i, (nhan, gia_tri, kieu) in enumerate(muc, 1):
        ghi_o(ws, i, 1, nhan, dam=True)
        ghi_o(ws, i, 2, gia_tri, kieu)
    hang = len(muc) + 2
    ghi_o(ws, hang, 1, "Chú giải ký hiệu", dam=True)
    for chu in CHU_GIAI:
        hang += 1
        ghi_o(ws, hang, 1, chu)
    for chu in (
        "Số liệu phản ánh dữ liệu đã ghi nhận trong hệ thống tại thời điểm trích xuất.",
        "Mẫu nghiệp vụ nội bộ — không phải biểu mẫu pháp lý.",
        *ghi_chu_them,
    ):
        hang += 2
        ghi_o(ws, hang, 1, chu)
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 70
    _thiet_lap_in(ws, dinh_nghia, tt, ngang=False, so_cot=2, hang_cuoi=hang)


def tao_so_chi_tiet(ket_qua, tt):
    """Workbook cho SO-01, SO-02, SO-03. Trả về (workbook, số dòng dữ liệu)."""
    dn, nhom = ket_qua.dinh_nghia, ket_qua.dinh_nghia.nhom
    wb = Workbook()
    ws = wb.active
    ws.title = dn.ma_gon
    cot = cac_cot(nhom, tt.du_lieu_ca_nhan_day_du)
    cot_in = so_cot_in(cot)
    them = []
    if nhom == Luot.Nhom.KHACH:
        them.append("Dữ liệu cá nhân: " + ("hiển thị đầy đủ" if tt.du_lieu_ca_nhan_day_du else "số giấy tờ đã che"))
    them.append(f"Vùng in gồm {cot_in} cột đầu; các cột sau chỉ dùng để tra cứu trong Excel.")
    hang_tieu_de = _ghi_dau_so(ws, dn, ket_qua.ky, ket_qua.chot, ket_qua.bo_loc, tt, cot_in, them)
    if ket_qua.dong:
        hang = _ghi_bang(ws, hang_tieu_de, cot, ket_qua.dong)
    else:
        _ghi_bang(ws, hang_tieu_de, cot, [])
        hang = hang_tieu_de + 1
        ghi_o(ws, hang, 1, "Không có lượt nào trong kỳ theo điều kiện đã chọn.")

    hang += 2
    ghi_o(ws, hang, 1, f"TỔNG HỢP SỐ LIỆU (tại thời điểm chốt {_dt(ket_qua.chot)})", dam=True)
    ct = ket_qua.chi_tieu
    for nhan, gia_tri, don_vi in _dong_chi_tieu(ct, nhom):
        hang += 1
        _ghi_cap(ws, hang, nhan, gia_tri, don_vi=don_vi) if gia_tri is not None else ghi_o(ws, hang, 2, nhan)
    hang += 1
    ghi_o(ws, hang, 2, f"Cân đối: Từ kỳ trước + Phát sinh = Hoàn tất + Chưa hoàn tất "
                       f"({ct['tu_ky_truoc']} + {ct['phat_sinh']} = {ct['hoan_tat']} + {ct['chua_hoan_tat']})")
    if ket_qua.bo_loc.co_loc:
        hang += 1
        ghi_o(ws, hang, 2, "Số liệu theo điều kiện lọc đã nêu ở đầu sổ.")
    hang += 1
    ghi_o(ws, hang, 2, "Chú giải ký hiệu và phụ lục đính chính: xem sheet Thong_tin_trich_xuat và Dinh_chinh.")
    hang = _ghi_khoi_ky(ws, hang + 2, tt, max(cot_in - 5, 4))

    _thiet_lap_in(ws, dn, tt, ngang=True, hang_tieu_de=hang_tieu_de, so_cot=cot_in, hang_cuoi=hang)
    _ghi_sheet_dinh_chinh(wb, ket_qua, tt)
    so_dong = len(ket_qua.dong)
    _ghi_sheet_thong_tin(wb, dn, ket_qua.ky, ket_qua.chot, ket_qua.bo_loc, tt, so_dong)
    return wb, so_dong


# ---------------------------------------------------------------------------
# Sổ tổng hợp
# ---------------------------------------------------------------------------

_TEN_NHOM = {Luot.Nhom.CBCS: "CBCS", Luot.Nhom.KHACH: "Khách", Luot.Nhom.PHUONG_TIEN: "Phương tiện"}


def _dau_phan(ws, ket_qua, tt, tieu_de, so_cot, ghi_chu=()):
    hang = _ghi_dau_so(ws, ket_qua.dinh_nghia, ket_qua.ky, ket_qua.chot, ket_qua.bo_loc, tt, so_cot,
                       ["Nhóm đối tượng: " + ", ".join(_TEN_NHOM[n] for n in ket_qua.nhom),
                        "Phân bổ: theo ngày lịch"])
    ghi_o(ws, hang, 1, tieu_de, dam=True)
    for chu in ghi_chu:
        hang += 1
        o = ghi_o(ws, hang, 1, chu)
        o.font = Font(italic=True)
    return hang + 2


def tao_so_tong_hop(ket_qua, tt):
    """Workbook cho SO-04. Trả về (workbook, số dòng dữ liệu của mọi phần)."""
    dn = ket_qua.dinh_nghia
    wb = Workbook()
    so_dong = {}

    # Phần A — không có dòng tổng cộng giữa các nhóm (BR-TONG-04).
    ws = wb.active
    ws.title = "A_Chi_tieu"
    cot_a = [
        Cot("Nhóm đối tượng", lambda m: _TEN_NHOM[m[0]], rong_toi_thieu=14),
        Cot("SỐ LƯỢT — Từ kỳ trước (lượt)", lambda m: m[1]["tu_ky_truoc"], "so", 12, 14),
        Cot("SỐ LƯỢT — Phát sinh trong kỳ (lượt)", lambda m: m[1]["phat_sinh"], "so", 12, 14),
        Cot("SỐ LƯỢT — Hoàn tất trong kỳ (lượt)", lambda m: m[1]["hoan_tat"], "so", 12, 14),
        Cot("SỐ LƯỢT — Chưa hoàn tất tại thời điểm chốt (lượt)", lambda m: m[1]["chua_hoan_tat"], "so", 14, 16),
        Cot("SỐ LƯỢT — Trong đó quá hạn (lượt)", lambda m: m[1]["qua_han"], "so", 12, 14),
        Cot("SỐ LƯỢT — Không định danh (lượt)", lambda m: "—" if m[1]["khong_dinh_danh"] is None
            else m[1]["khong_dinh_danh"], "so", 12, 14),
        Cot("SỐ LƯỢT — Có đính chính (lượt)", lambda m: m[1]["luot_dinh_chinh"], "so", 12, 14),
        Cot("SỐ ĐỐI TƯỢNG — Duy nhất, có định danh (đối tượng)", lambda m: m[1]["doi_tuong_duy_nhat"], "so", 14, 16),
        Cot("Bản ghi đã huỷ, không tính (bản ghi)", lambda m: m[1]["ban_ghi_huy"], "so", 12, 14),
        Cot("Đối chiếu sổ riêng", lambda m: {"CBCS": "SO-01", "KHACH": "SO-02", "PT": "SO-03"}[m[0]]),
    ]
    hang = _dau_phan(ws, ket_qua, tt, "A. CHỈ TIÊU THEO NHÓM ĐỐI TƯỢNG", len(cot_a),
                     ["Số liệu của từng nhóm đối tượng được tính riêng và không cộng gộp giữa các nhóm."])
    dong_a = [(n, ket_qua.phan_a[n]) for n in ket_qua.nhom]
    hang_cuoi = _ghi_bang_hon_hop(ws, hang, cot_a, dong_a)
    ghi_o(ws, hang_cuoi + 1, 1, "(Không có dòng tổng cộng giữa các nhóm.)")
    hang_cuoi = _ghi_khoi_ky(ws, hang_cuoi + 3, tt, max(len(cot_a) - 4, 4))
    _thiet_lap_in(ws, dn, tt, ngang=False, hang_tieu_de=hang, so_cot=len(cot_a), hang_cuoi=hang_cuoi)
    so_dong["A"] = len(dong_a)

    # Phần B — theo ngày (BR-TONG-07).
    ws = wb.create_sheet("B_Theo_ngay")
    cot_b = [Cot("Ngày", lambda m: m[0], "ngay", 12)]
    nhan_b = {
        Luot.Nhom.CBCS: ("CBCS — lượt phát sinh (lượt)", "CBCS — có hoạt động (đối tượng)"),
        Luot.Nhom.KHACH: ("Khách — lượt phát sinh (lượt)", "Khách có định danh — có hoạt động (đối tượng)"),
        Luot.Nhom.PHUONG_TIEN: ("Phương tiện — lượt phát sinh (lượt)", "Phương tiện có biển — có hoạt động (đối tượng)"),
    }
    for n in ket_qua.nhom:
        cot_b.append(Cot(nhan_b[n][0], lambda m, n=n: m[1][n][0], "so", 12, 16))
        cot_b.append(Cot(nhan_b[n][1], lambda m, n=n: m[1][n][1], "so", 12, 18))
    hang = _dau_phan(ws, ket_qua, tt, "B. PHÂN BỔ THEO NGÀY", len(cot_b), [
        "Cột “có hoạt động” đếm đối tượng có ít nhất một lượt trong ngày; một đối tượng có thể có mặt ở "
        "nhiều ngày nên dòng cuối là số đối tượng duy nhất của cả kỳ, không phải tổng các ngày.",
    ])
    hang_cuoi = _ghi_bang(ws, hang, cot_b, ket_qua.phan_b)
    hang_cuoi += 1
    ghi_o(ws, hang_cuoi, 1, "Cộng / Duy nhất cả kỳ (cột đối tượng không phải tổng các dòng)", dam=True, khung=True)
    for i, n in enumerate(ket_qua.nhom):
        tong, duy_nhat = ket_qua.phan_b_cuoi[n]
        ghi_o(ws, hang_cuoi, 2 + 2 * i, tong, "so", dam=True, khung=True)
        ghi_o(ws, hang_cuoi, 3 + 2 * i, duy_nhat, "so", dam=True, khung=True)
    _thiet_lap_in(ws, dn, tt, ngang=False, hang_tieu_de=hang, so_cot=len(cot_b), hang_cuoi=hang_cuoi)
    so_dong["B"] = len(ket_qua.phan_b)

    # Phần C — theo cổng, đếm sự kiện (BR-TONG-08).
    ws = wb.create_sheet("C_Theo_cong")
    cot_c = [Cot("Cổng", lambda m: m[0].ten, rong_toi_thieu=12)]
    for n in ket_qua.nhom:
        cot_c.append(Cot(f"{_TEN_NHOM[n]} — sự kiện Vào", lambda m, n=n: m[1][n][0], "so", 10, 14))
        cot_c.append(Cot(f"{_TEN_NHOM[n]} — sự kiện Ra", lambda m, n=n: m[1][n][1], "so", 10, 14))
    hang = _dau_phan(ws, ket_qua, tt, "C. PHÂN BỔ THEO CỔNG (số sự kiện)", len(cot_c))
    hang_cuoi = _ghi_bang(ws, hang, cot_c, ket_qua.phan_c)
    hang_cuoi += 1
    ghi_o(ws, hang_cuoi, 1, "Cộng", dam=True, khung=True)
    for i, n in enumerate(ket_qua.nhom):
        vao, ra = ket_qua.phan_c_cuoi[n]
        ghi_o(ws, hang_cuoi, 2 + 2 * i, vao, "so", dam=True, khung=True)
        ghi_o(ws, hang_cuoi, 3 + 2 * i, ra, "so", dam=True, khung=True)
    _thiet_lap_in(ws, dn, tt, ngang=False, hang_tieu_de=hang, so_cot=len(cot_c), hang_cuoi=hang_cuoi)
    so_dong["C"] = len(ket_qua.phan_c)

    # Phần D — lượt chưa hoàn tất, bảng con theo nhóm, không cộng chung.
    ws = wb.create_sheet("D_Chua_hoan_tat")
    dong_d = []
    for n in ket_qua.nhom:
        for stt, d in enumerate(ket_qua.phan_d[n], 1):
            dong_d.append((n, stt, d))

    def doi_tuong(m):
        n, _, d = m
        if not ket_qua.xem_chi_tiet[n]:
            return f"(không có quyền xem chi tiết — mã lượt {d.luot.pk})"
        l = d.luot
        if n == Luot.Nhom.CBCS:
            return f"{l.can_bo.ma} {l.can_bo.ho_ten}"
        if n == Luot.Nhom.KHACH:
            return l.khach.ho_ten + (f" (thẻ {l.so_the})" if l.so_the else "")
        return f"{l.phuong_tien.hien_thi_bien_so} {l.phuong_tien.loai}"

    def su_kien_mo(m):
        d = m[2]
        loai = "RA" if d.luot.chieu == Luot.Chieu.RA_VAO else "VAO"
        return f"{'Ra' if loai == 'RA' else 'Vào'} — {d.su_kien(loai).cong.ten}"

    cot_d = [
        Cot("Nhóm", lambda m: _TEN_NHOM[m[0]]),
        Cot("STT trong nhóm", lambda m: m[1], "so"),
        Cot("Đối tượng", doi_tuong, rong_toi_thieu=20, rong_toi_da=50),
        Cot("Sự kiện mở", su_kien_mo, rong_toi_thieu=12),
        Cot("Thời điểm mở", lambda m: m[2].luot.mo_thoi_diem, "ngay_gio", 19),
        Cot("Hạn dự kiến", lambda m: m[2].luot.han_du_kien, "ngay_gio", 19),
        Cot("Thời lượng tính đến chốt (tối thiểu)", lambda m: ket_qua.chot - m[2].luot.mo_thoi_diem,
            "thoi_luong", 12),
        Cot("Trạng thái", lambda m: m[2].nhan_trang_thai, rong_toi_thieu=14),
        Cot("Mã lượt", lambda m: m[2].luot.pk, "so"),
    ]
    tom_tat = [f"{_TEN_NHOM[n]}: {len(ket_qua.phan_d[n])} lượt" for n in ket_qua.nhom]
    hang = _dau_phan(ws, ket_qua, tt, "D. LƯỢT CHƯA HOÀN TẤT TẠI THỜI ĐIỂM CHỐT", len(cot_d),
                     ["Số lượng theo từng nhóm (không cộng chung): " + "; ".join(tom_tat)])
    hang_cuoi = _ghi_bang(ws, hang, cot_d, dong_d)
    if not dong_d:
        hang_cuoi += 1
        ghi_o(ws, hang_cuoi, 1, "Không có lượt chưa hoàn tất.")
    _thiet_lap_in(ws, dn, tt, ngang=True, hang_tieu_de=hang, so_cot=len(cot_d), hang_cuoi=hang_cuoi)
    so_dong["D"] = len(dong_d)

    # Phần E — tóm tắt đính chính, huỷ.
    ws = wb.create_sheet("E_Dinh_chinh_huy")
    cot_e = [
        Cot("Nhóm", lambda m: _TEN_NHOM[m]),
        Cot("Số lượt có đính chính (lượt)", lambda m: ket_qua.phan_a[m]["luot_dinh_chinh"], "so", 14),
        Cot("Số đính chính (lần)", lambda m: ket_qua.phan_a[m]["so_dinh_chinh"], "so", 12),
        Cot("Số bản ghi đã huỷ (bản ghi)", lambda m: ket_qua.phan_a[m]["ban_ghi_huy"], "so", 14),
        Cot("Xem chi tiết", lambda m: "Phụ lục đính chính của " +
            {"CBCS": "SO-01", "KHACH": "SO-02", "PT": "SO-03"}[m], rong_toi_thieu=20),
    ]
    hang = _dau_phan(ws, ket_qua, tt, "E. TÓM TẮT ĐÍNH CHÍNH VÀ HUỶ", len(cot_e))
    hang_cuoi = _ghi_bang_hon_hop(ws, hang, cot_e, ket_qua.nhom)
    _thiet_lap_in(ws, dn, tt, ngang=False, hang_tieu_de=hang, so_cot=len(cot_e), hang_cuoi=hang_cuoi)
    so_dong["E"] = len(ket_qua.nhom)

    tong = sum(so_dong.values())
    _ghi_sheet_thong_tin(wb, dn, ket_qua.ky, ket_qua.chot, ket_qua.bo_loc, tt, tong,
                         ["Số dòng theo phần: " + "; ".join(f"{k}: {v}" for k, v in so_dong.items())])
    return wb, so_dong


def _ghi_bang_hon_hop(ws, hang_dau, cac_cot, dong):
    """Bảng nhỏ, không bật bộ lọc Excel (không hữu ích với 3 dòng)."""
    hang = hang_dau
    rong = [do_dai_hien_thi(c.tieu_de, "text") // 2 for c in cac_cot]
    for j, c in enumerate(cac_cot, 1):
        o = ghi_o(ws, hang, j, c.tieu_de, dam=True, khung=True, xuong_dong=True)
        o.fill = _NEN_TIEU_DE
        o.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for d in dong:
        hang += 1
        for j, c in enumerate(cac_cot, 1):
            gia_tri = c.lay(d)
            kieu = "text" if isinstance(gia_tri, str) else c.kieu
            ghi_o(ws, hang, j, gia_tri, kieu, khung=True)
            rong[j - 1] = max(rong[j - 1], do_dai_hien_thi(gia_tri, kieu))
    for j, c in enumerate(cac_cot, 1):
        ws.column_dimensions[get_column_letter(j)].width = min(max(rong[j - 1] + 2, c.rong_toi_thieu), c.rong_toi_da)
    ws.freeze_panes = ws.cell(row=hang_dau + 1, column=1)
    return hang


def luu(wb):
    bo_dem = io.BytesIO()
    wb.save(bo_dem)
    return bo_dem.getvalue()
